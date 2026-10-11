from openpyxl import load_workbook

from app.models import DataRequest
from app.routers.property_data import PROPERTY_SHEETS, _property_workbook


def test_property_template_contains_complete_v5_sheets():
    workbook = load_workbook(_property_workbook(), data_only=True)
    assert workbook.sheetnames == [
        "Instructions", "QuickImport", "Developers", "Projects", "Configurations",
        "Towers", "Units", "NearbyLocations", "Contacts", "LegalCompliance",
        "PricingCharges", "MediaDocuments", "SourcesVerification", "Lists",
        "Examples", "ImportMap",
    ]
    assert workbook["QuickImport"]["A1"].value == "Developer / Builder"
    assert workbook["Projects"]["A1"].value == "id"
    assert workbook["Configurations"]["A1"].value == "id"
    assert workbook["Towers"]["A1"].value == "id"
    assert "unit_number" not in [cell.value for cell in workbook["Projects"][1]]
    assert "floor" not in [cell.value for cell in workbook["Projects"][1]]
    assert set(PROPERTY_SHEETS) == {"Developers", "Projects", "Towers", "Units", "NearbyLocations"}


def test_property_data_routes_are_registered_in_production_app():
    from app.main_portal import app
    paths = {route.path for route in app.routes}
    assert "/admin/property-data" in paths
    assert "/admin/property-data/template" in paths
    assert "/admin/property-data/import" in paths
    assert "/admin/property-data/save" in paths
    assert "/admin/property-data/export" in paths
    assert "/admin/property-data/export/request" in paths
    assert "/admin/property-data/export/download/{request_id}" in paths
    assert "/super-admin/property-data" in paths
    assert "/super-admin/property-data/template" in paths
    assert "/super-admin/property-data/load-godrej-nest" in paths
    assert "/super-admin/property-data/import" in paths
    assert "/super-admin/property-data/save" in paths
    assert "/super-admin/property-data-export-requests" in paths
    assert "/super-admin/property-data-export-requests/{request_id}/approve" in paths
    assert "/super-admin/property-data-export-requests/{request_id}/reject" in paths
    assert "/super-admin/property-data-notifications" in paths


def test_property_data_ui_has_role_aware_import_export_and_online_edit_controls():
    from pathlib import Path
    template = Path("app/templates/property_data.html").read_text(encoding="utf-8")
    assert "property_data_base" in template
    assert 'action="{{ property_data_base }}/load-godrej-nest"' in template
    assert 'action="{{ property_data_base }}/import"' in template
    assert 'action="{{ property_data_base }}/save"' in template
    assert "/admin/property-data/export/request" in template
    assert "/super-admin/property-data/export" in template
    assert "/super-admin/property-data-export-requests" in template
    assert "property_data_base }}/export/download/" in template
    assert "csrf_token" in template
    assert "Edit" in template
    assert "Download Complete V5 Excel Template" in template


def test_super_admin_property_data_path_is_explicitly_bridged_to_admin_handler():
    from app.portal_boundary import _SUPER_ADMIN_LEGACY_ROUTES, _super_admin_redirect_for_legacy_admin_path
    assert ("/admin/property-data", "/super-admin/property-data") in _SUPER_ADMIN_LEGACY_ROUTES
    assert _super_admin_redirect_for_legacy_admin_path("/admin/property-data") == "/super-admin/property-data"
    assert _super_admin_redirect_for_legacy_admin_path("/admin/property-data/template") == "/super-admin/property-data/template"
    assert _super_admin_redirect_for_legacy_admin_path("/admin/property-data/import") == "/super-admin/property-data/import"
    assert _super_admin_redirect_for_legacy_admin_path("/admin/property-data/save") == "/super-admin/property-data/save"


def test_property_export_uses_existing_data_request_model():
    assert DataRequest.__tablename__ == "data_requests"
    assert {column.name for column in DataRequest.__table__.columns} >= {
        "requester_id", "action", "entity_type", "filters_json", "reason", "status", "decided_by_id", "decided_at"
    }
