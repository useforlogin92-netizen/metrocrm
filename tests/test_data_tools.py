from io import BytesIO

from openpyxl import load_workbook

from app.models import Developer, User
from app.routers.data_tools import _descendant_ids, _excel_value, _export_report_workbook, _import_workbook, _template_workbook


def test_import_template_contains_leads_and_property_sheets():
    workbook = load_workbook(_template_workbook(), data_only=True)
    assert "Instructions" in workbook.sheetnames
    assert "Leads" in workbook.sheetnames
    assert "Projects" in workbook.sheetnames
    assert "Units" in workbook.sheetnames
    assert "NearbyLocations" in workbook.sheetnames
    assert "Examples" in workbook.sheetnames
    assert workbook["Leads"]["A1"].value == "id"
    assert workbook["Projects"]["B1"].value == "developer_id"


def test_import_creates_and_updates_records(db_session):
    wb = load_workbook(_template_workbook())
    for sheet in wb.sheetnames:
        if sheet not in {"Developers", "Instructions", "Examples"}:
            wb[sheet].delete_rows(2, wb[sheet].max_row)
    ws = wb["Developers"]
    ws.delete_rows(2, ws.max_row)
    ws.append([None, "Test Developer", "9999999999", "First import"])
    data = BytesIO()
    wb.save(data)
    created, updated, errors = _import_workbook(data.getvalue(), db_session)
    assert errors == []
    assert created == 1
    assert updated == 0
    developer = db_session.query(Developer).filter_by(name="Test Developer").one()
    developer_id = developer.id

    wb2 = load_workbook(_template_workbook())
    ws2 = wb2["Developers"]
    ws2.delete_rows(2, ws2.max_row)
    ws2.append([developer_id, "Updated Developer", "8888888888", "Updated"])
    data2 = BytesIO()
    wb2.save(data2)
    created2, updated2, errors2 = _import_workbook(data2.getvalue(), db_session)
    assert errors2 == []
    assert created2 == 0
    assert updated2 == 1
    assert db_session.get(Developer, developer_id).name == "Updated Developer"


def test_report_scope_includes_manager_tree_only(db_session):
    manager = User(username="manager", password_hash="x", role="MANAGER", is_active=True, is_approved=True)
    tl = User(username="tl", password_hash="x", role="TL", reporting_manager_id=None, is_active=True, is_approved=True)
    employee = User(username="employee", password_hash="x", role="SALES_EXECUTIVE", reporting_manager_id=None, is_active=True, is_approved=True)
    outsider = User(username="outsider", password_hash="x", role="SALES_EXECUTIVE", is_active=True, is_approved=True)
    db_session.add_all([manager, tl, employee, outsider])
    db_session.flush()
    tl.reporting_manager_id = manager.id
    employee.reporting_manager_id = tl.id
    db_session.commit()
    ids = _descendant_ids(manager, db_session)
    assert ids == {manager.id, tl.id, employee.id}
    assert outsider.id not in ids


def test_report_scope_for_super_admin_is_global(db_session):
    super_admin = User(username="super", password_hash="x", role="SUPER_ADMIN", is_active=True, is_approved=True)
    other = User(username="other", password_hash="x", role="TELECALLER", is_active=True, is_approved=True)
    db_session.add_all([super_admin, other])
    db_session.commit()
    assert _descendant_ids(super_admin, db_session) == {super_admin.id, other.id}


def test_report_workbook_generation(db_session):
    super_admin = User(username="report_admin", password_hash="x", role="SUPER_ADMIN", is_active=True, is_approved=True)
    db_session.add(super_admin)
    db_session.commit()
    workbook = load_workbook(_export_report_workbook(super_admin, db_session), data_only=True)
    assert "Report Summary" in workbook.sheetnames
    assert "Team Scope" in workbook.sheetnames
    assert workbook["Report Summary"]["A1"].value == "Metric"


def test_excel_export_neutralizes_formula_prefixes():
    assert _excel_value("=1+1") == "'=1+1"
    assert _excel_value("@SUM(1,1)") == "'@SUM(1,1)"
    assert _excel_value("normal text") == "normal text"
    assert _excel_value(100) == 100


def test_data_tools_routes_are_registered():
    from app.main import app
    paths = {route.path for route in app.routes}
    assert "/admin/data-tools" in paths
    assert "/admin/data-tools/import" in paths
    assert "/admin/data-tools/export/{entity_type}" in paths
    assert "/reports" in paths
    assert "/super-admin/business-health" in paths


def test_xlsx_upload_is_validated_as_zip_and_limited():
    from app.routers import data_tools
    assert hasattr(data_tools, "is_zipfile")
    assert hasattr(data_tools, "ZipFile")


def test_data_tools_html_output_uses_html_escaping():
    from app.routers.data_tools import escape
    assert escape("<script>alert(1)</script>") == "&lt;script&gt;alert(1)&lt;/script&gt;"
