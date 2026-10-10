from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models import Customer, Lead, User
from app.routers.lead_import import LEAD_FIELDS, _import_leads, _lead_template


def _actor(db_session):
    actor = User(username="lead-import-template-test", password_hash="x", role="ADMIN", is_active=True, is_approved=True)
    db_session.add(actor)
    db_session.commit()
    return actor


def test_template_has_blank_import_sheet_source_dropdown_and_separate_example():
    workbook = load_workbook(_lead_template())
    leads = workbook["Leads"]
    example = workbook["Example"]
    assert [cell.value for cell in leads[1]] == LEAD_FIELDS
    assert all(cell.value is None for cell in leads[2])
    assert example["D2"].value == "Rahul Sharma"
    assert leads.data_validations.dataValidation
    source_column = LEAD_FIELDS.index("source") + 1
    expected_letter = leads.cell(row=1, column=source_column).column_letter
    validation = leads.data_validations.dataValidation[0]
    assert validation.formula1.startswith('"Magicbricks,Meta Ads,99acres')
    assert f"{expected_letter}2" in validation.sqref


def test_import_accepts_legacy_template_without_serial_column(db_session):
    actor = _actor(db_session)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    sheet.append(LEAD_FIELDS[1:])
    output = BytesIO()
    workbook.save(output)

    created, updated, errors = _import_leads(output.getvalue(), actor, db_session)

    assert (created, updated, errors) == (0, 0, [])
    assert db_session.query(Lead).count() == 0


def test_serial_number_only_row_does_not_create_customer_or_lead(db_session):
    actor = _actor(db_session)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"
    sheet.append(LEAD_FIELDS)
    sheet.append([1] + [None] * (len(LEAD_FIELDS) - 1))
    output = BytesIO()
    workbook.save(output)

    created, updated, errors = _import_leads(output.getvalue(), actor, db_session)

    assert (created, updated, errors) == (0, 0, [])
    assert db_session.query(Lead).count() == 0
    assert db_session.query(Customer).count() == 0
