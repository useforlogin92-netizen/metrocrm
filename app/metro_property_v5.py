"""Metro Properties V5 property-data workbook and structured import helpers.

Metro-only module.  It keeps the existing Property Master/WhatsApp flow intact
while adding the complete V5 workbook and a guided data-entry contract.
"""
from __future__ import annotations

import json
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

V5_SHEETS = (
    "Instructions", "QuickImport", "Developers", "Projects", "Configurations",
    "Towers", "Units", "NearbyLocations", "Contacts", "LegalCompliance",
    "PricingCharges", "MediaDocuments", "SourcesVerification", "Lists",
    "Examples", "ImportMap",
)

HEADERS = {
    "QuickImport": [
        "Developer / Builder","Project Name","City","State","Locality / Area","Address",
        "Property Type","BHK / Configuration","Carpet Area (sq.ft.)","Built-up Area (sq.ft.)",
        "Rate / sq.ft.","Bundle / Total Price","Parking","Furnishing","Tower / Wing",
        "Unit / Flat No.","Floor","Total Floors","Availability","Project Status",
        "Possession Date","MahaRERA Number","RERA Project URL","Developer Contact Name",
        "Developer Contact Number","Developer Email","Website","Nearby Station",
        "Station Distance","Nearby Location / Landmark","Configuration Notes",
        "Payment Plan / Charges","CP Commission","Bank / Loan Details",
        "Title Deed / Legal Notes","Photos URL","Floor Plan URL","Brochure URL",
        "Video URL","Data Source","Source Reference Date","Last Verified Date","Remarks",
    ],
    "Developers": [
        "id","name","legal_company_name","brand_name","contact_name","contact_number",
        "whatsapp_number","email","website","office_address","city","state","pincode",
        "rera_promoter_details","company_gst","developer_status","details","data_source",
        "source_url","source_reference_date","last_verified_date","remarks",
    ],
    "Projects": [
        "id","developer_id","name","project_aliases","state","district","city","region","zone",
        "locality","sub_locality","sector","address","pincode","latitude","longitude",
        "nearby_station","station_distance","nearby_landmarks","rera","rera_project_url",
        "possession_date","builder_commitment_date","launch_date","construction_status",
        "project_status","land_area","total_towers","total_floors","total_units","phase",
        "maintenance_estimate","society_name","oc_cc_status","title_conveyance_status",
        "amenities_json","photos_json","brochure","video_url","website","description",
        "payment_plan_charges","bank_loan_details","data_source","source_url",
        "source_reference_date","last_verified_date","remarks",
    ],
    "Configurations": [
        "id","project_id","developer_id","bhk_configuration","carpet_area_sqft",
        "built_up_area_sqft","super_built_up_area_sqft","rate_per_sqft","basic_price",
        "total_price","floor_rise_charge","plc_charge","parking_charge","club_charge",
        "infrastructure_charge","development_charge","legal_charge","maintenance_charge",
        "corpus_charge","other_charges","stamp_duty","registration","gst","all_in_price",
        "parking","furnishing","availability","tower_wing","status","verification_status",
        "data_source","source_url","source_reference_date","last_verified_date","notes",
    ],
    "Towers": [
        "id","project_id","name","wing","floors","units_per_floor","lifts","service_lift",
        "fire_lift","parking_levels","podium_basement","possession_date","construction_status",
        "facing_view","details","data_source","source_url","last_verified_date",
    ],
    "Units": [
        "id","project_id","tower_id","configuration_id","unit_number","floor",
        "total_building_floors","stack","bhk","property_type","listing_type","transaction_type",
        "new_or_resale","carpet_area","built_up_area","super_built_up_area","balcony_area",
        "terrace_area","deck_area","dry_balcony_area","plot_area","ceiling_height",
        "entrance_direction","door_direction","window_direction","facing","view","corner_unit",
        "road_view","garden_view","sea_view","pool_view","open_view","hill_view","price",
        "price_per_sqft","basic_price","total_price","parking","parking_count","parking_type",
        "parking_number","parking_level","parking_charge","furnishing","furnishing_details_json",
        "possession","rera","availability","status","latitude","longitude","amenities_json",
        "photos_json","floor_plan","brochure","video_url","description","payment_plan_charges",
        "cp_commission","bank_loan_details","title_deed_legal_notes","developer_contact_name",
        "developer_contact_number","developer_email","data_source","source_url",
        "source_reference_date","last_verified_date","remarks","other_details",
    ],
    "NearbyLocations": [
        "id","project_id","category","name","sub_category","distance","distance_unit",
        "travel_time","transport_mode","details","latitude","longitude","verified",
        "verified_at","data_source","source_url",
    ],
    "Contacts": [
        "id","entity_type","entity_id","contact_type","name","designation","mobile",
        "whatsapp","email","alternate_mobile","office_phone","notes","source_url",
        "last_verified_date",
    ],
    "LegalCompliance": [
        "id","project_id","unit_id","rera_number","rera_project_url","rera_status",
        "promoter_name","registered_project_name","phase","possession_date","oc_status",
        "cc_status","title_status","conveyance_status","litigation_flag","document_type",
        "document_reference","verification_status","verified_by","verified_date","source_url",
        "legal_notes",
    ],
    "PricingCharges": [
        "id","project_id","configuration_id","unit_id","price_type","quoted_price",
        "negotiated_price","rate_per_sqft","floor_rise","plc","parking","club",
        "infrastructure","development","legal","maintenance","corpus","other_charges",
        "stamp_duty","registration","gst","all_in_price","payment_plan","effective_date",
        "expiry_date","source_url","verified_date","remarks",
    ],
    "MediaDocuments": [
        "id","entity_type","entity_id","media_type","title","url","file_reference",
        "document_number","uploaded_by","uploaded_date","validity_date","verification_status",
        "verified_by","verified_date","version","source_url","remarks",
    ],
    "SourcesVerification": [
        "id","entity_type","entity_id","field_name","source_type","source_name","source_url",
        "source_reference_date","confidence","verification_status","conflict_flag",
        "conflict_details","checked_by","checked_date","last_updated",
    ],
}

LISTS = {
    "Property Type":["Residential","Commercial","Plot","Office","Shop","Warehouse","Industrial","Land","Villa","Row House","Bungalow","Other"],
    "BHK / Configuration":["1 RK","1 BHK","1.5 BHK","2 BHK","2.5 BHK","3 BHK","3.5 BHK","4 BHK","4.5 BHK","5 BHK","5+ BHK","Studio","Duplex","Triplex","Penthouse","Jodi / Combination","Villa","Row House","Bungalow","Other"],
    "Availability":["Available","Limited","YES","NO","HOLD","Booked","Sold","Blocked","Coming Soon","Unknown","On Request"],
    "Project Status":["ACTIVE","UPCOMING","UNDER CONSTRUCTION","COMPLETED","ON HOLD","CANCELLED"],
    "Furnishing":["Unfurnished","Semi Furnished","Furnished","Part Furnished","Bare Shell","To Verify"],
    "Parking":["No Parking","1 Car","2 Cars","3 Cars","4+ Cars","Car + Bike","2 Car + Bike","Open","Covered","Podium","Basement","Mechanical","Stack","EV","To Verify"],
    "Verification Status":["UNVERIFIED","PUBLIC_SOURCE","ADMIN_VERIFIED","NEEDS_REVIEW","CONFLICT"],
    "Construction Status":["PLANNED","LAUNCHED","UNDER CONSTRUCTION","NEAR COMPLETION","COMPLETED","ON HOLD","CANCELLED"],
    "Source Type":["WhatsApp","Excel","Manual","Website","RERA","Developer","Channel Partner","AI Research","Portal","Other"],
    "Nearby Category":["Railway Station","Metro Station","Bus","School","College","Hospital","Mall","Market","Bank","ATM","Restaurant","Office District","Airport","Highway","Expressway","Landmark","Other"],
}

def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()

def _num(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return None

def _json(v: Any, default=None):
    if v in (None, ""):
        return default
    if isinstance(v, (dict, list)):
        return v
    try:
        return json.loads(str(v))
    except (TypeError, ValueError):
        return default

def _set_if(obj: Any, attr: str, value: Any):
    if hasattr(obj, attr) and value not in (None, ""):
        setattr(obj, attr, value)

def _to_date(v: Any):
    if isinstance(v, datetime):
        return v
    if v in (None, ""):
        return None
    text = _text(v)
    for fmt in ("%Y-%m-%d","%d-%m-%Y","%d/%m/%Y","%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None

def _style_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_column:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(1, ws.max_row)}"
    for i, cell in enumerate(ws[1], 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(cell.value or ""))+3, 14), 34)

def build_v5_workbook():
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Instructions")
    rows = [
        ["Metro Properties – Complete Property Master Import Template",""],
        ["Version","V5.0 – merged Quick Import + structured Property Master"],
        ["Purpose","Single approved workbook for Quick Import and Developer → Project → Configuration → Tower/Wing → Unit → Nearby Location data."],
        ["Core rule","Unknown information must remain blank / To Verify. Never invent property information."],
        ["Import order","Developers → Projects → Configurations → Towers → Units → NearbyLocations → supporting sheets."],
        ["Update rule","Blank id = create. Existing id = update that exact record."],
        ["Configuration rule","One row = one Project + one BHK/Configuration + one Carpet Area. Multiple carpet-area rows are allowed for the same BHK."],
        ["Source rule","Retain Data Source, Source URL/reference date and Last Verified Date wherever available."],
        ["Review rule","Imported records pass validation, duplicate/conflict checks and approval before becoming trusted Property Master records."],
        ["Compatibility","QuickImport retains the original 43-field Property Master import format; structured sheets add the richer master details."],
        ["JSON rule","Fields ending _json must contain valid JSON arrays/objects."],
        ["Date rule","Prefer YYYY-MM-DD; timestamps may use YYYY-MM-DD HH:MM:SS."],
    ]
    for row in rows: ws.append(row)
    ws.column_dimensions["A"].width=30; ws.column_dimensions["B"].width=120
    ws["A1"].font=Font(bold=True,size=14,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="17365D")
    ws["B1"].fill=PatternFill("solid",fgColor="17365D")

    for name, headers in HEADERS.items():
        sh=wb.create_sheet(name); sh.append(headers); _style_sheet(sh)

    ls=wb.create_sheet("Lists")
    for col,(name,values) in enumerate(LISTS.items(),1):
        ls.cell(1,col,name)
        for row,value in enumerate(values,2): ls.cell(row,col,value)
    _style_sheet(ls)

    ex=wb.create_sheet("Examples")
    for row in [
        ["Sheet","Field","Example Value"],
        ["Configurations","bhk_configuration","2 BHK"],
        ["Configurations","carpet_area_sqft",620],
        ["Configurations","rate_per_sqft",26451],
        ["Configurations","total_price",16400000],
        ["Configurations","availability","Available"],
        ["Configurations","parking","1 Car"],
        ["Units","tower_id","B-WING"],
        ["Units","unit_number","1204"],
        ["NearbyLocations","category","Railway Station"],
        ["NearbyLocations","name","Kandivali Station"],
        ["NearbyLocations","distance",1.8],
        ["NearbyLocations","distance_unit","km"],
        ["Projects","amenities_json",'["Gym","Swimming Pool","Club House"]'],
    ]: ex.append(row)
    _style_sheet(ex)

    im=wb.create_sheet("ImportMap")
    im.append(["Original 43-field column","V5 destination","Rule"])
    mapping = {
        "Developer / Builder":"Developers.name / Projects.developer_id",
        "Project Name":"Projects.name",
        "City":"Projects.city","State":"Projects.state","Locality / Area":"Projects.locality",
        "Address":"Projects.address","Property Type":"Units.property_type",
        "BHK / Configuration":"Configurations.bhk_configuration / Units.bhk",
        "Carpet Area (sq.ft.)":"Configurations.carpet_area_sqft / Units.carpet_area",
        "Built-up Area (sq.ft.)":"Configurations.built_up_area_sqft / Units.built_up_area",
        "Rate / sq.ft.":"Configurations.rate_per_sqft / Units.price_per_sqft",
        "Bundle / Total Price":"Configurations.total_price / Units.total_price",
        "Parking":"Configurations.parking / Units.parking",
        "Furnishing":"Configurations.furnishing / Units.furnishing",
        "Tower / Wing":"Towers.name / Units.tower_id","Unit / Flat No.":"Units.unit_number",
        "Floor":"Units.floor","Total Floors":"Towers.floors / Units.total_building_floors",
        "Availability":"Configurations.availability / Units.availability",
        "Project Status":"Projects.project_status","Possession Date":"Projects.possession_date / Units.possession",
        "MahaRERA Number":"Projects.rera / LegalCompliance.rera_number",
        "RERA Project URL":"Projects.rera_project_url","Developer Contact Name":"Contacts.name",
        "Developer Contact Number":"Contacts.mobile","Developer Email":"Contacts.email",
        "Website":"Projects.website","Nearby Station":"NearbyLocations.name","Station Distance":"NearbyLocations.distance",
        "Nearby Location / Landmark":"NearbyLocations.name","Configuration Notes":"Configurations.notes",
        "Payment Plan / Charges":"Projects.payment_plan_charges / PricingCharges.payment_plan",
        "CP Commission":"Units.cp_commission","Bank / Loan Details":"Projects.bank_loan_details / Units.bank_loan_details",
        "Title Deed / Legal Notes":"Units.title_deed_legal_notes / LegalCompliance.legal_notes",
        "Photos URL":"MediaDocuments.url / Units.photos_json","Floor Plan URL":"MediaDocuments.url / Units.floor_plan",
        "Brochure URL":"MediaDocuments.url / Units.brochure","Video URL":"MediaDocuments.url / Units.video_url",
        "Data Source":"SourcesVerification.source_type","Source Reference Date":"SourcesVerification.source_reference_date",
        "Last Verified Date":"SourcesVerification.last_verified_date","Remarks":"Units.remarks / Projects.remarks",
    }
    for h in HEADERS["QuickImport"]:
        im.append([h,mapping.get(h,""),"Map during normalization; do not silently discard information."])
    _style_sheet(im)
    return wb

def workbook_bytes():
    wb=build_v5_workbook()
    out=BytesIO()
    wb.save(out); out.seek(0)
    return out

def _row_map(ws):
    headers=[_text(c.value) for c in ws[1]]
    return [dict(zip(headers,[row[i] if i < len(row) else None for i in range(len(headers))])) for row in ws.iter_rows(min_row=2,values_only=True)]

def _upsert_developer(db, Developer, raw, cache):
    name=_text(raw.get("name") or raw.get("Developer / Builder"))
    if not name: return None
    key=name.casefold()
    obj=None
    raw_id=_text(raw.get("id"))
    if raw_id.isdigit(): obj=db.get(Developer,int(raw_id))
    if obj is None: obj=cache.get(key)
    if obj is None: obj=db.query(Developer).filter(Developer.name.ilike(name)).first()
    if obj is None:
        obj=Developer(name=name); db.add(obj)
        db.flush()
    elif obj.id is None:
        db.flush()
    _set_if(obj,"contact",raw.get("contact") or raw.get("contact_number") or raw.get("Developer Contact Number"))
    details={k:v for k,v in raw.items() if k not in {"id","name"} and v not in (None,"")}
    obj.details=json.dumps(details,ensure_ascii=False,default=str)
    cache[key]=obj
    return obj

def _project_payload(raw):
    return {
        "state":raw.get("state") or raw.get("State"),"district":raw.get("district") or raw.get("District"),"city":raw.get("city") or raw.get("City"),
        "locality":raw.get("locality") or raw.get("Locality / Area"),"sector":raw.get("sector"),
        "address":raw.get("address") or raw.get("Address"),"pincode":raw.get("pincode"),
        "latitude":_num(raw.get("latitude")),"longitude":_num(raw.get("longitude")),
        "rera":raw.get("rera") or raw.get("MahaRERA Number"),
        "possession_date":_to_date(raw.get("possession_date") or raw.get("Possession Date")),
        "builder_commitment_date":_to_date(raw.get("builder_commitment_date")),
        "construction_status":raw.get("construction_status") or raw.get("Project Status"),
        "amenities_json":json.dumps(_json(raw.get("amenities_json"),[]) or [],ensure_ascii=False),
        "photos_json":json.dumps(_json(raw.get("photos_json"),[]) or [],ensure_ascii=False),
        "brochure":raw.get("brochure") or raw.get("Brochure URL"),
        "description":raw.get("description") or "",
    }

def import_v5_workbook(db, models, wb, actor):
    """Import the V5 workbook. Returns (created, updated, warnings)."""
    Developer, Project, PropertyConfiguration, Tower, Unit, NearbyLocation = models
    created=updated=0; warnings=[]
    developer_cache={}; project_cache={}; tower_cache={}; config_cache={}

    # QuickImport is deliberately normalized into the same structured records.
    # It is the primary user-facing upload sheet: one row = one project + one
    # BHK/configuration + one carpet-area combination. Re-uploading the same
    # natural key updates the existing configuration instead of creating a duplicate.
    if "QuickImport" in wb.sheetnames:
        for row_no, raw in enumerate(_row_map(wb["QuickImport"]), start=2):
            if not any(v not in (None, "") for v in raw.values()):
                continue
            pname=_text(raw.get("Project Name"))
            if not pname:
                warnings.append(f"QuickImport row {row_no} skipped: Project Name is required")
                continue
            dev_name=_text(raw.get("Developer / Builder"))
            if not dev_name:
                warnings.append(f"QuickImport row {row_no} skipped: Developer / Builder is required")
                continue
            bhk=_text(raw.get("BHK / Configuration"))
            if not bhk:
                warnings.append(f"QuickImport row {row_no} skipped: BHK / Configuration is required")
                continue

            dev=_upsert_developer(db,Developer,{
                "name":dev_name,
                "contact":raw.get("Developer Contact Number"),
                "email":raw.get("Developer Email"),
                "website":raw.get("Website"),
                "contact_name":raw.get("Developer Contact Name"),
            },developer_cache)
            key=(dev.id or id(dev),pname.casefold())
            project=project_cache.get(key)
            if project is None:
                project=db.query(Project).filter(
                    Project.developer_id==dev.id,Project.name.ilike(pname)
                ).first()
            if project is None:
                project=Project(developer_id=dev.id,name=pname)
                db.add(project)
                created+=1
            else:
                updated+=1

            prior_description=project.description
            for attr,val in _project_payload(raw).items():
                _set_if(project,attr,val)
            extra={k:v for k,v in raw.items() if v not in (None,"")}
            try:
                existing_desc=json.loads(prior_description or "{}")
                if not isinstance(existing_desc,dict):
                    existing_desc={}
            except (TypeError,ValueError):
                existing_desc={}
            rows=existing_desc.get("quick_import_rows")
            if not isinstance(rows,list):
                rows=[]
            # Keep the latest value for the same natural row key.
            area=_num(raw.get("Carpet Area (sq.ft.)"))
            tower_name=_text(raw.get("Tower / Wing"))
            unit_number=_text(raw.get("Unit / Flat No."))
            row_key="|".join([pname.casefold(),bhk.casefold(),str(area or ""),tower_name.casefold(),unit_number.casefold()])
            rows=[x for x in rows if not isinstance(x,dict) or x.get("_natural_key")!=row_key]
            extra["_natural_key"]=row_key
            rows.append(extra)
            existing_desc["quick_import_rows"]=rows[-250:]
            project.description=json.dumps(existing_desc,ensure_ascii=False,default=str)
            project_cache[key]=project
            db.flush()

            cfg=db.query(PropertyConfiguration).filter(
                PropertyConfiguration.project_id==project.id,
                PropertyConfiguration.configuration==bhk,
                PropertyConfiguration.carpet_area==area,
            ).first()
            if cfg is None:
                cfg=PropertyConfiguration(project_id=project.id,configuration=bhk,carpet_area=area)
                db.add(cfg)
            cfg.rate_per_sqft=_num(raw.get("Rate / sq.ft."))
            cfg.price_including_all=_num(raw.get("Bundle / Total Price"))
            cfg.availability=_text(raw.get("Availability")) or None
            cfg.source_details=json.dumps(extra,ensure_ascii=False,default=str)
            db.flush()

            tower_name=_text(raw.get("Tower / Wing"))
            tower=None
            if tower_name:
                tower=db.query(Tower).filter(
                    Tower.project_id==project.id,Tower.name.ilike(tower_name)
                ).first()
                if tower is None:
                    tower=Tower(project_id=project.id,name=tower_name)
                    db.add(tower)
                    db.flush()
                tower.details=json.dumps({
                    "total_floors":raw.get("Total Floors"),
                    "project_status":raw.get("Project Status"),
                    "source":raw.get("Data Source"),
                },ensure_ascii=False,default=str)
                possession=_to_date(raw.get("Possession Date"))
                if possession:
                    tower.possession_date=possession

            unit_number=_text(raw.get("Unit / Flat No."))
            if unit_number:
                if tower:
                    unit=db.query(Unit).filter(
                        Unit.project_id==project.id,
                        Unit.unit_number==unit_number,
                        Unit.tower_id==tower.id,
                    ).first()
                    if unit is None:
                        unassigned=db.query(Unit).filter(
                            Unit.project_id==project.id,
                            Unit.unit_number==unit_number,
                            Unit.tower_id.is_(None),
                        ).all()
                        if len(unassigned)==1:
                            unit=unassigned[0]
                        else:
                            unit=Unit(project_id=project.id,unit_number=unit_number)
                            db.add(unit)
                else:
                    matches=db.query(Unit).filter(
                        Unit.project_id==project.id,Unit.unit_number==unit_number
                    ).all()
                    if len(matches)>1:
                        warnings.append(
                            f"QuickImport row {row_no}: unit {unit_number!r} is ambiguous across towers; unit update skipped"
                        )
                        unit=None
                    elif matches:
                        unit=matches[0]
                    else:
                        unit=Unit(project_id=project.id,unit_number=unit_number)
                        db.add(unit)
                if unit is not None:
                    if tower:
                        unit.tower_id=tower.id
                    unit.bhk=bhk
                    unit.property_type=_text(raw.get("Property Type")) or None
                    unit.carpet_area=area
                    unit.built_up_area=_num(raw.get("Built-up Area (sq.ft.)"))
                    unit.price=cfg.price_including_all
                    unit.price_per_sqft=cfg.rate_per_sqft
                    unit.parking=_text(raw.get("Parking")) or None
                    unit.furnishing=_text(raw.get("Furnishing")) or None
                    unit.availability=_text(raw.get("Availability")) or None
                    unit.status=(_text(raw.get("Availability")) or "AVAILABLE").upper()
                    floor=_num(raw.get("Floor"))
                    if floor is not None:
                        unit.floor=int(floor)
                    total_floors=_num(raw.get("Total Floors"))
                    if total_floors is not None:
                        unit.total_building_floors=int(total_floors)
                    unit.possession=_text(raw.get("Possession Date")) or None
                    unit.rera=_text(raw.get("MahaRERA Number")) or None
                    unit.floor_plan=_text(raw.get("Floor Plan URL")) or None
                    unit.brochure=_text(raw.get("Brochure URL")) or None
                    unit.other_details=json.dumps(extra,ensure_ascii=False,default=str)

            nearby_pairs=[
                ("Railway Station",_text(raw.get("Nearby Station")),_text(raw.get("Station Distance"))),
                ("Landmark",_text(raw.get("Nearby Location / Landmark")),""),
            ]
            for category,name,distance in nearby_pairs:
                if not name:
                    continue
                nearby=db.query(NearbyLocation).filter(
                    NearbyLocation.project_id==project.id,
                    NearbyLocation.category.ilike(category),
                    NearbyLocation.name.ilike(name),
                ).first()
                if nearby is None:
                    nearby=NearbyLocation(project_id=project.id,category=category,name=name)
                    db.add(nearby)
                nearby.distance=distance or nearby.distance
                nearby.details=json.dumps({"source":_text(raw.get("Data Source"))},ensure_ascii=False)

    # Structured sheets: master rows first.
    if "Developers" in wb.sheetnames:
        for raw in _row_map(wb["Developers"]):
            if any(v not in (None,"") for v in raw.values()): _upsert_developer(db,Developer,raw,developer_cache)
        db.flush()

    if "Projects" in wb.sheetnames:
        for raw in _row_map(wb["Projects"]):
            if not any(v not in (None,"") for v in raw.values()): continue
            name=_text(raw.get("name"))
            if not name: warnings.append("Projects row skipped: name is required"); continue
            dev=None
            did=_text(raw.get("developer_id"))
            if did.isdigit(): dev=db.get(Developer,int(did))
            if dev is None:
                # Accept a developer name as a convenience even though the canonical field is developer_id.
                dev=_upsert_developer(db,Developer,{"name":raw.get("developer") or raw.get("developer_name")},developer_cache)
            if dev is None: warnings.append(f"Project '{name}' skipped: developer_id/developer name is required"); continue
            obj=None; rid=_text(raw.get("id"))
            if rid.isdigit(): obj=db.get(Project,int(rid))
            if obj is None: obj=db.query(Project).filter(Project.developer_id==dev.id,Project.name.ilike(name)).first()
            if obj is None: obj=Project(developer_id=dev.id,name=name); db.add(obj); created+=1
            else: updated+=1
            prior_description=obj.description
            for attr,val in _project_payload(raw).items(): _set_if(obj,attr,val)
            try:
                description_data=json.loads(prior_description or "{}")
                if not isinstance(description_data,dict):
                    description_data={}
            except (TypeError,ValueError):
                description_data={"legacy_description":prior_description} if prior_description else {}
            description_data.update({k:v for k,v in raw.items() if v not in (None,"")})
            obj.description=json.dumps(description_data,ensure_ascii=False,default=str)
            db.flush(); project_cache[name.casefold()]=obj

    # Configurations
    if "Configurations" in wb.sheetnames:
        for raw in _row_map(wb["Configurations"]):
            if not any(v not in (None,"") for v in raw.values()): continue
            project=None; pid=_text(raw.get("project_id"))
            if pid.isdigit(): project=db.get(Project,int(pid))
            if project is None:
                pname=_text(raw.get("project_name")); project=project_cache.get(pname.casefold()) if pname else None
            if project is None: warnings.append("Configuration row skipped: project_id/project_name not found"); continue
            bhk=_text(raw.get("bhk_configuration"))
            if not bhk: warnings.append(f"Configuration for project #{project.id} skipped: bhk_configuration is required"); continue
            area=_num(raw.get("carpet_area_sqft")); key=(project.id,bhk.casefold(),area)
            obj=None; rid=_text(raw.get("id"))
            if rid.isdigit(): obj=db.get(PropertyConfiguration,int(rid))
            if obj is None: obj=db.query(PropertyConfiguration).filter(PropertyConfiguration.project_id==project.id,PropertyConfiguration.configuration==bhk,PropertyConfiguration.carpet_area==area).first()
            if obj is None: obj=PropertyConfiguration(project_id=project.id,configuration=bhk,carpet_area=area); db.add(obj); created+=1
            else: updated+=1
            for attr in ("rate_per_sqft","price_including_all","availability","negotiation"):
                _set_if(obj,attr,raw.get(attr))
            obj.source_details=json.dumps({k:v for k,v in raw.items() if v not in (None,"")},ensure_ascii=False,default=str)
            db.flush(); config_cache[(project.id,bhk.casefold(),area)]=obj

    # Towers
    if "Towers" in wb.sheetnames:
        for raw in _row_map(wb["Towers"]):
            if not any(v not in (None,"") for v in raw.values()): continue
            pid=_text(raw.get("project_id")); project=db.get(Project,int(pid)) if pid.isdigit() else None
            if project is None: continue
            name=_text(raw.get("name")); wing=_text(raw.get("wing")); full=name or wing
            if not full: warnings.append(f"Tower for project #{project.id} skipped: name is required"); continue
            obj=None; rid=_text(raw.get("id"))
            if rid.isdigit(): obj=db.get(Tower,int(rid))
            if obj is None: obj=db.query(Tower).filter(Tower.project_id==project.id,Tower.name.ilike(full)).first()
            if obj is None: obj=Tower(project_id=project.id,name=full); db.add(obj); created+=1
            else: updated+=1
            if raw.get("floors") not in (None,""): _set_if(obj,"floors",int(float(raw["floors"])))
            _set_if(obj,"possession_date",_to_date(raw.get("possession_date")))
            obj.details=json.dumps({k:v for k,v in raw.items() if v not in (None,"")},ensure_ascii=False,default=str)
            db.flush(); tower_cache[(project.id,full.casefold())]=obj

    # Units
    if "Units" in wb.sheetnames:
        unit_attrs=("unit_number","floor","total_building_floors","bhk","property_type","carpet_area","built_up_area","super_built_up_area","plot_area","price","price_per_sqft","parking","facing","furnishing","possession","rera","availability","status","latitude","longitude","floor_plan","brochure","description","other_details")
        numeric={"floor","total_building_floors","carpet_area","built_up_area","super_built_up_area","plot_area","price","price_per_sqft","latitude","longitude"}
        for raw in _row_map(wb["Units"]):
            if not any(v not in (None,"") for v in raw.values()): continue
            pid=_text(raw.get("project_id")); project=db.get(Project,int(pid)) if pid.isdigit() else None
            if project is None: continue
            unit_no=_text(raw.get("unit_number"))
            if not unit_no: warnings.append(f"Unit for project #{project.id} skipped: unit_number is required"); continue
            tid=_text(raw.get("tower_id")); tower=db.get(Tower,int(tid)) if tid.isdigit() else None
            cid=_text(raw.get("configuration_id")); cfg=db.get(PropertyConfiguration,int(cid)) if cid.isdigit() else None
            if tower:
                obj=db.query(Unit).filter(
                    Unit.project_id==project.id,
                    Unit.unit_number==unit_no,
                    Unit.tower_id==tower.id,
                ).first()
                if obj is None:
                    unassigned=db.query(Unit).filter(
                        Unit.project_id==project.id,
                        Unit.unit_number==unit_no,
                        Unit.tower_id.is_(None),
                    ).all()
                    if len(unassigned)==1:
                        obj=unassigned[0]
                    else:
                        obj=Unit(project_id=project.id,unit_number=unit_no)
                        db.add(obj); created+=1
                else:
                    updated+=1
            else:
                matches=db.query(Unit).filter(
                    Unit.project_id==project.id,Unit.unit_number==unit_no
                ).all()
                if len(matches)>1:
                    warnings.append(f"Unit row for project #{project.id}: unit {unit_no!r} is ambiguous across towers; update skipped")
                    continue
                obj=matches[0] if matches else None
                if obj is None:
                    obj=Unit(project_id=project.id,unit_number=unit_no)
                    db.add(obj); created+=1
                else:
                    updated+=1
            if tower:
                obj.tower_id=tower.id
            for attr in unit_attrs:
                val=raw.get(attr)
                if attr in numeric: val=_num(val)
                _set_if(obj,attr,val)
            if cfg: pass
            obj.other_details=json.dumps({k:v for k,v in raw.items() if k not in set(unit_attrs) and v not in (None,"")},ensure_ascii=False,default=str)
            db.flush()

    if "NearbyLocations" in wb.sheetnames:
        for raw in _row_map(wb["NearbyLocations"]):
            if not any(v not in (None,"") for v in raw.values()): continue
            pid=_text(raw.get("project_id")); project=db.get(Project,int(pid)) if pid.isdigit() else None
            if project is None: continue
            name=_text(raw.get("name")); category=_text(raw.get("category"))
            if not name or not category: continue
            obj=db.query(NearbyLocation).filter(NearbyLocation.project_id==project.id,NearbyLocation.name.ilike(name),NearbyLocation.category.ilike(category)).first()
            if obj is None: obj=NearbyLocation(project_id=project.id,category=category,name=name); db.add(obj); created+=1
            else: updated+=1
            _set_if(obj,"distance",raw.get("distance"))
            obj.details=json.dumps({k:v for k,v in raw.items() if v not in (None,"")},ensure_ascii=False,default=str)

    return created, updated, warnings
