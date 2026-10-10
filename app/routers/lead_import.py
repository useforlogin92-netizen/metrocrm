from __future__ import annotations

from io import BytesIO
from datetime import datetime
import json

from fastapi import APIRouter, Depends, File, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.csrf import csrf_token, require_csrf
from app.database import get_db
from app.models import AuditEvent, Customer, Lead, User

router = APIRouter(tags=["offline-lead-import"])
templates = Jinja2Templates(directory="app/templates")

LEAD_FIELDS = [
    "sr_no", "lead_id", "customer_id", "customer_name", "mobile", "alternate_mobile", "email",
    "source", "status", "assigned_user_id", "original_owner_id", "requirement",
    "budget_min", "budget_max", "preferred_location", "bhk", "property_type", "buy_rent",
    "customer_notes",
]


def _user(request: Request, db: Session):
    from app.main import _current_user
    return _current_user(request, db)


def _guard(request: Request, db: Session):
    user = _user(request, db)
    return user if user and user.role in {"SUPER_ADMIN", "ADMIN"} else None


def _lead_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    rows = [
        ["MetroCRM Offline Lead Upload Template"],
        ["Purpose", "Fill this Excel file offline and upload it to MetroCRM."],
        ["Rule 1", "Keep the Leads sheet column names unchanged."],
        ["Rule 2", "Leave lead_id blank for a new lead. Use an existing lead_id to update that lead."],
        ["Rule 3", "customer_id is optional. If blank, MetroCRM will find a customer by mobile/email or create one from customer_name."],
        ["Rule 4", "For a new customer, customer_name and mobile/email should be supplied."],
        ["Rule 5", "assigned_user_id must be a valid MetroCRM user ID."],
        ["Rule 6", "The complete workbook is validated before any changes are committed."],
        ["Rule 7", "Only Super Admin and Admin can upload offline leads."],
        ["Rule 8", "Admin imports do not require approval, but Super Admin receives an activity notice."],
        ["Rule 9", "Sr No is for spreadsheet reference only; MetroCRM generates its own lead ID."],
        ["Rule 10", "Only rows on the Leads sheet are imported. The Example sheet is for reference and is never imported."],
        ["Rule 11", "Use the Source dropdown for Magicbricks, Meta Ads, 99acres, Website, Referral, Walk-in, WhatsApp, Instagram, Facebook, Housing.com, NoBroker, Other, or Offline Import."],
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 105
    ws["A1"].font = Font(bold=True, size=16)

    sh = wb.create_sheet("Leads")
    sh.append(LEAD_FIELDS)
    examples = {
        "sr_no": 1,
        "lead_id": "",
        "customer_id": "",
        "customer_name": "Rahul Sharma",
        "mobile": "9876543210",
        "alternate_mobile": "9123456780",
        "email": "rahul@example.com",
        "source": "Magicbricks",
        "status": "NEW",
        "assigned_user_id": 1,
        "original_owner_id": 1,
        "requirement": "2 BHK near Thane West",
        "budget_min": 5000000,
        "budget_max": 8000000,
        "preferred_location": "Thane West",
        "bhk": "2 BHK",
        "property_type": "APARTMENT",
        "buy_rent": "BUY",
        "customer_notes": "Interested in ready possession",
    }
    # Keep the import sheet blank so uploading the untouched template cannot
    # accidentally create a sample customer/lead in the production CRM.
    example_sheet = wb.create_sheet("Example")
    example_sheet.append(LEAD_FIELDS)
    example_sheet.append([examples.get(x, "") for x in LEAD_FIELDS])
    for cell in example_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    example_sheet.freeze_panes = "A2"
    for idx, field in enumerate(LEAD_FIELDS, 1):
        example_sheet.column_dimensions[get_column_letter(idx)].width = min(30, max(13, len(field) + 4))
    for cell in sh[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    sh.freeze_panes = "A2"
    last_column = get_column_letter(len(LEAD_FIELDS))
    sh.auto_filter.ref = f"A1:{last_column}5000"
    source_letter = get_column_letter(LEAD_FIELDS.index("source") + 1)
    source_validation = DataValidation(
        type="list",
        formula1='"Magicbricks,Meta Ads,99acres,Website,Referral,Walk-in,WhatsApp,Instagram,Facebook,Housing.com,NoBroker,Other,Offline Import"',
        allow_blank=False,
    )
    source_validation.error = "Select a source from the dropdown."
    source_validation.errorTitle = "Invalid lead source"
    source_validation.prompt = "Choose where this lead came from."
    source_validation.promptTitle = "Lead source"
    sh.add_data_validation(source_validation)
    source_validation.add(f"{source_letter}2:{source_letter}5000")
    for idx, field in enumerate(LEAD_FIELDS, 1):
        sh.column_dimensions[get_column_letter(idx)].width = min(30, max(13, len(field) + 4))

    ex = wb.create_sheet("Field Guide")
    ex.append(["Field", "Required?", "What to enter"])
    guide = {
        "sr_no": ("No", "Spreadsheet serial number for reference only; ignored by MetroCRM"),
        "lead_id": ("No", "Existing lead ID for update; blank creates a lead"),
        "customer_id": ("No", "Existing customer ID if known"),
        "customer_name": ("For new customer", "Customer full name"),
        "mobile": ("Recommended", "Primary customer mobile number"),
        "source": ("Yes", "Lead source: Magicbricks, Meta Ads, 99acres, Website, Referral, Walk-in, WhatsApp, Instagram, Facebook, Housing.com, NoBroker, Other, or Offline Import"),
        "status": ("Yes", "NEW, CONTACTED, QUALIFIED, etc."),
        "assigned_user_id": ("No", "MetroCRM user ID responsible for the lead"),
        "requirement": ("No", "Free-text property requirement"),
        "budget_min": ("No", "Minimum budget in INR"),
        "budget_max": ("No", "Maximum budget in INR"),
        "preferred_location": ("No", "Preferred area/city/locality"),
        "bhk": ("No", "Example: 2 BHK"),
        "property_type": ("No", "Example: APARTMENT"),
        "buy_rent": ("No", "BUY or RENT"),
        "customer_notes": ("No", "Customer notes"),
    }
    for field in LEAD_FIELDS:
        req, desc = guide.get(field, ("No", "Optional"))
        ex.append([field, req, desc])
    for cell in ex[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ex.freeze_panes = "A2"
    ex.column_dimensions["A"].width = 30
    ex.column_dimensions["B"].width = 20
    ex.column_dimensions["C"].width = 90

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _as_decimal(value):
    if value in (None, ""):
        return None
    return float(value)


def _find_or_create_customer(db: Session, raw: dict, existing_customer_id: int | None = None):
    if existing_customer_id:
        customer = db.get(Customer, existing_customer_id)
        if customer is None:
            raise ValueError(f"customer_id {existing_customer_id} does not exist")
    else:
        customer = None
    mobile = str(raw.get("mobile") or "").strip() or None
    email = str(raw.get("email") or "").strip().lower() or None
    name = str(raw.get("customer_name") or "").strip()
    if customer is None and (mobile or email):
        filters = []
        if mobile:
            filters.append(Customer.mobile == mobile)
        if email:
            filters.append(Customer.email == email)
        customer = db.scalar(select(Customer).where(or_(*filters)).order_by(Customer.id.asc())) if filters else None
    if customer is None:
        if not name:
            raise ValueError("customer_name is required when a new customer must be created")
        customer = Customer(name=name, mobile=mobile, alternate_mobile=str(raw.get("alternate_mobile") or "").strip() or None, email=email, notes=str(raw.get("customer_notes") or "").strip() or None)
        db.add(customer)
        db.flush()
    else:
        if name:
            customer.name = name
        if mobile:
            customer.mobile = mobile
        alternate = str(raw.get("alternate_mobile") or "").strip()
        if alternate:
            customer.alternate_mobile = alternate
        if email:
            customer.email = email
        notes = str(raw.get("customer_notes") or "").strip()
        if notes:
            customer.notes = notes
    return customer


def _import_leads(content: bytes, actor: User, db: Session):
    wb = load_workbook(BytesIO(content), data_only=True)
    if "Leads" not in wb.sheetnames:
        return 0, 0, ["Workbook must contain a Leads sheet"]
    ws = wb["Leads"]
    headers = [str(x.value).strip() if x.value is not None else "" for x in ws[1]]
    has_serial_column = bool(headers and headers[0].lower() == "sr_no")
    normalized_headers = headers[1:] if has_serial_column else headers
    expected_fields = LEAD_FIELDS[1:]
    if normalized_headers != expected_fields:
        return 0, 0, ["Leads sheet columns do not match the MetroCRM template"]
    users = {u.id for u in db.scalars(select(User)).all()}
    created = updated = 0
    errors: list[str] = []
    try:
        for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(v not in (None, "") for v in values):
                continue
            if has_serial_column:
                values = values[1:]
            raw = {expected_fields[i]: values[i] for i in range(len(expected_fields))}
            try:
                lead_id = int(raw["lead_id"]) if raw.get("lead_id") not in (None, "") else None
                lead = db.get(Lead, lead_id) if lead_id else Lead()
                if lead_id and lead is None:
                    raise ValueError(f"lead_id {lead_id} does not exist")
                assigned = raw.get("assigned_user_id")
                owner = raw.get("original_owner_id")
                if assigned not in (None, "") and int(assigned) not in users:
                    raise ValueError(f"assigned_user_id {assigned} does not exist")
                if owner not in (None, "") and int(owner) not in users:
                    raise ValueError(f"original_owner_id {owner} does not exist")
                customer_id = int(raw["customer_id"]) if raw.get("customer_id") not in (None, "") else None
                customer = _find_or_create_customer(db, raw, customer_id)
                lead.customer_id = customer.id
                lead.source = str(raw.get("source") or "OFFLINE_IMPORT").strip()
                lead.status = str(raw.get("status") or "NEW").strip()
                lead.assigned_user_id = int(assigned) if assigned not in (None, "") else None
                lead.original_owner_id = int(owner) if owner not in (None, "") else None
                lead.requirement = str(raw.get("requirement") or "").strip() or None
                lead.budget_min = _as_decimal(raw.get("budget_min"))
                lead.budget_max = _as_decimal(raw.get("budget_max"))
                lead.preferred_location = str(raw.get("preferred_location") or "").strip() or None
                lead.bhk = str(raw.get("bhk") or "").strip() or None
                lead.property_type = str(raw.get("property_type") or "").strip() or None
                lead.buy_rent = str(raw.get("buy_rent") or "").strip() or None
                db.add(lead)
                if lead_id:
                    updated += 1
                else:
                    created += 1
            except Exception as exc:
                errors.append(f"Leads row {row_no}: {exc}")
                break
        if errors:
            db.rollback()
        else:
            db.flush()
            db.add(AuditEvent(actor_user_id=actor.id, action="LEADS_IMPORTED", entity_type="Lead", entity_id=None, details=json.dumps({"created": created, "updated": updated, "mode": "offline_excel"})))
            db.commit()
    except Exception as exc:
        db.rollback()
        errors.append(f"Lead import validation failed: {exc}")
    return created, updated, errors


@router.get("/super-admin/leads/offline-import", response_class=HTMLResponse)
@router.get("/admin/leads/offline-import", response_class=HTMLResponse)
def offline_lead_import_page(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/dashboard", status_code=303)
    return templates.TemplateResponse("lead_offline_import.html", {"request": request, "username": actor.username, "role": actor.role, "avatar": actor.username[:1].upper(), "csrf_token": csrf_token(request), "message": request.query_params.get("message"), "error": request.query_params.get("error"), "show_user_management": True, "show_integrations": False, "show_alerts": False, "show_approvals": True})


@router.get("/super-admin/leads/offline-import/template")
@router.get("/admin/leads/offline-import/template")
def offline_lead_template(request: Request, db: Session = Depends(get_db)):
    if not _guard(request, db):
        return RedirectResponse("/dashboard", status_code=303)
    return StreamingResponse(_lead_template(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=MetroCRM_Offline_Lead_Upload_Template.xlsx"})


@router.post("/super-admin/leads/offline-import", dependencies=[Depends(require_csrf)])
@router.post("/admin/leads/offline-import", dependencies=[Depends(require_csrf)])
async def offline_lead_import(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/dashboard", status_code=303)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return RedirectResponse("/admin/leads/offline-import?error=Only+.xlsx+files+are+supported", status_code=303)
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        return RedirectResponse("/admin/leads/offline-import?error=Maximum+file+size+is+10+MB", status_code=303)
    created, updated, errors = _import_leads(content, actor, db)
    if errors:
        return RedirectResponse("/admin/leads/offline-import?error=" + str(errors[0]).replace(" ", "+")[:350], status_code=303)
    return RedirectResponse(f"/admin/leads/offline-import?message=Imported+{created}+new+and+updated+{updated}+leads", status_code=303)


@router.get("/super-admin/lead-import-notifications")
def lead_import_notifications(request: Request, db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor or actor.role != "SUPER_ADMIN":
        return {"items": []}
    events = db.scalars(select(AuditEvent).where(AuditEvent.action.in_(["LEADS_IMPORTED"])).order_by(AuditEvent.id.desc()).limit(12)).all()
    users = {u.id: u.username for u in db.scalars(select(User).where(User.id.in_({e.actor_user_id for e in events if e.actor_user_id}))).all()} if events else {}
    items = []
    for event in events:
        try:
            details = json.loads(event.details or "{}")
        except ValueError:
            details = {}
        items.append({"message": f"Admin lead import: {details.get('created', 0)} new, {details.get('updated', 0)} updated", "username": users.get(event.actor_user_id, "Unknown"), "created_at": event.created_at.isoformat() if event.created_at else ""})
    return {"items": items}
