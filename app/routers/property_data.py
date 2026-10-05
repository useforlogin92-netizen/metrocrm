from __future__ import annotations

import json
import re
from datetime import date, datetime, timezone
from urllib.parse import quote_plus
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, File, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from sqlalchemy import String, cast, or_, select
from sqlalchemy.orm import Session

from app.csrf import csrf_token, require_csrf
from app.database import get_db
from app.models import AuditEvent, Booking, Commission, DataRequest, Developer, NearbyLocation, Project, PropertyConfiguration, SiteVisit, Token, Tower, Unit, User
from app.routers.data_tools import _convert_excel_value, _export_workbook, _import_workbook, _models, _save_workbook, _template_workbook
from app.property_master import PropertyUpdateSource
from app.services.property_update_engine import apply_property_master
from app.metro_property_v5 import workbook_bytes, import_v5_workbook, _num, _to_date
from tools.seed_godrej_nest_test import seed_godrej_nest_test

router = APIRouter(tags=["property-data"])
templates = Jinja2Templates(directory="app/templates")
PROPERTY_SHEETS = ("Developers", "Projects", "Towers", "Units", "NearbyLocations")
PROPERTY_MODELS = {name: _models()[name] for name in PROPERTY_SHEETS}


def _user(request: Request, db: Session):
    from app.main import _current_user
    return _current_user(request, db)


def _guard(request: Request, db: Session):
    user = _user(request, db)
    if not user or user.role not in {"SUPER_ADMIN", "ADMIN"}:
        return None
    return user


def _property_workbook() -> BytesIO:
    """Return the Metro Properties V5 complete workbook."""
    return workbook_bytes()

def _split_source_values(value) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [part.strip() for part in re.split(r"\s*/\s*", text) if part.strip()]


def _parse_source_area(value):
    text = str(value or "").strip().replace(",", "")
    match = re.search(r"\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def _parse_source_price(value) -> tuple[float | None, float | None]:
    text = str(value or "").strip().lower().replace(",", "")
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        return None, None
    number = float(match.group())
    if re.search(r"\b(?:cr|crore)\b", text):
        return number * 10000000, None
    if re.search(r"\b(?:lakh|lac)\b", text):
        return number * 100000, None
    if "psf" in text or re.search(r"\d+(?:\.\d+)?\s*k\b", text):
        return None, number * 1000 if "k" in text else number
    return None, None


def _source_configuration_rows(raw_config: dict, availability: str | None) -> list[dict]:
    configuration = str(raw_config.get("Configuration") or raw_config.get("BHK") or "").strip()
    if not configuration:
        return []
    areas = _split_source_values(raw_config.get("Carpet Area"))
    prices = _split_source_values(raw_config.get("Bundle Price"))
    if len(areas) == len(prices) and areas:
        pairs = list(zip(areas, prices))
    elif len(areas) > 1:
        pairs = [(area, prices[0] if len(prices) == 1 else None) for area in areas]
    elif len(prices) > 1:
        pairs = [(areas[0] if areas else None, None) for _ in prices]
    else:
        pairs = [(areas[0] if areas else None, prices[0] if prices else None)]
    rows = []
    for area_text, price_text in pairs:
        price, rate = _parse_source_price(price_text)
        rows.append({
            "configuration": configuration,
            "carpet_area": _parse_source_area(area_text),
            "price_including_all": price,
            "rate_per_sqft": rate,
            "negotiation": "YES" if "negotiat" in str(price_text or "").lower() else None,
            "availability": str(availability or "").strip() or None,
            "_source_details": raw_config,
        })
    return rows


def _source_tower_rows(configurations: list[dict]) -> list[dict]:
    seen = set()
    towers = []
    for raw in configurations:
        structure = str(raw.get("Tower/Wing / Structure") or "").strip()
        if not structure or structure in seen:
            continue
        seen.add(structure)
        towers.append({
            "tower_name": structure,
            "total_floors": None,
            "possession_date": raw.get("Possession"),
        })
    return towers

def _display_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


@router.get("/admin/property-data", response_class=HTMLResponse)
def property_data_page(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    selected = request.query_params.get("entity", "Projects")
    if selected not in PROPERTY_MODELS:
        selected = "Projects"
    model, fields = PROPERTY_MODELS[selected]
    search = request.query_params.get("search", "").strip()[:120]
    if search:
        pattern = f"%{search}%"
        if selected == "Projects":
            project_columns = [
                Project.id, Project.name, Project.state, Project.district, Project.city,
                Project.locality, Project.sector, Project.address, Project.pincode,
                Project.latitude, Project.longitude, Project.rera, Project.possession_date,
                Project.builder_commitment_date, Project.construction_status,
                Project.amenities_json, Project.photos_json, Project.brochure, Project.description,
            ]
            unit_columns = [
                Unit.id, Unit.unit_number, Unit.bhk, Unit.property_type, Unit.carpet_area,
                Unit.built_up_area, Unit.super_built_up_area, Unit.plot_area, Unit.price,
                Unit.price_per_sqft, Unit.parking, Unit.facing, Unit.furnishing, Unit.possession,
                Unit.rera, Unit.availability, Unit.status, Unit.amenities_json,
                Unit.photos_json, Unit.floor_plan, Unit.brochure, Unit.description, Unit.other_details,
            ]
            conditions = [cast(column, String).ilike(pattern) for column in project_columns]
            conditions.append(cast(Developer.name, String).ilike(pattern))
            unit_match = select(Unit.id).where(
                Unit.project_id == Project.id,
                or_(*[cast(column, String).ilike(pattern) for column in unit_columns]),
            ).exists()
            conditions.append(unit_match)
            stmt = select(Project).outerjoin(Developer, Developer.id == Project.developer_id).where(or_(*conditions))
        else:
            columns = [column for column in model.__table__.columns]
            stmt = select(model).where(or_(*[cast(column, String).ilike(pattern) for column in columns]))
        rows = db.scalars(stmt.order_by(model.id.desc()).limit(100)).all()
    else:
        rows = db.scalars(select(model).order_by(model.id.desc()).limit(100)).all()
    editing_id = request.query_params.get("id")
    edit_row = db.get(model, int(editing_id)) if editing_id and editing_id.isdigit() else None
    export_requests = db.scalars(
        select(DataRequest)
        .where(DataRequest.requester_id == actor.id, DataRequest.action == "EXPORT", DataRequest.entity_type == "PropertyData")
        .order_by(DataRequest.id.desc()).limit(10)
    ).all()
    return templates.TemplateResponse("property_data.html", {
        "request": request, "username": actor.username, "role": actor.role, "avatar": actor.username[:1].upper(),
        "csrf_token": csrf_token(request), "entities": PROPERTY_SHEETS, "selected": selected, "fields": fields,
        "rows": rows, "edit_row": edit_row, "display_value": _display_value, "getattr": getattr, "search": search,
        "export_requests": export_requests,
        "message": request.query_params.get("message"), "error": request.query_params.get("error"),
        "show_user_management": actor.role in {"SUPER_ADMIN", "ADMIN"}, "show_integrations": False,
        "show_alerts": False, "show_approvals": actor.role in {"SUPER_ADMIN", "ADMIN"},
    })


@router.get("/admin/property-data/v5-entry", response_class=HTMLResponse)
def property_data_v5_entry(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    project_id = request.query_params.get("project_id", "").strip()
    values = {}
    if project_id.isdigit():
        project = db.get(Project, int(project_id))
        if project:
            developer = db.get(Developer, project.developer_id)
            values = {
                "developer": developer.name if developer else "",
                "project_name": project.name,
                "state": project.state or "",
                "district": project.district or "",
                "locality": project.locality or "",
                "pincode": project.pincode or "",
                "address": project.address or "",
            }
    return templates.TemplateResponse("property_data_v5_entry.html", {
        "request": request,
        "username": actor.username,
        "role": actor.role,
        "avatar": actor.username[:1].upper(),
        "csrf_token": csrf_token(request),
        "project_id": project_id,
        "values": values,
        "base": "/super-admin" if actor.role == "SUPER_ADMIN" else "/admin",
        "error": request.query_params.get("error"),
        "show_user_management": actor.role in {"SUPER_ADMIN", "ADMIN"},
        "show_integrations": False,
        "show_alerts": False,
        "show_approvals": actor.role in {"SUPER_ADMIN", "ADMIN"},
    })


@router.post("/admin/property-data/v5-save", dependencies=[Depends(require_csrf)])
async def property_data_v5_save(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)

    form = await request.form()
    def val(name, default=""):
        value = form.get(name, default)
        return str(value).strip() if value is not None else default
    def vals(name):
        return [str(x).strip() for x in form.getlist(name) if str(x).strip()]

    try:
        developer_name = val("developer")
        project_name = val("project_name")
        if not developer_name or not project_name:
            raise ValueError("Developer / Builder and Project Name are required")

        developer = db.query(Developer).filter(Developer.name.ilike(developer_name)).first()
        if developer is None:
            developer = Developer(name=developer_name)
            db.add(developer)
            db.flush()
        contact_parts = []
        for key,label in (("developer_contact","Contact"),("developer_whatsapp","WhatsApp"),("developer_email","Email")):
            x=val(key)
            if x: contact_parts.append(f"{label}: {x}")
        if contact_parts:
            developer.contact=" | ".join(contact_parts)

        project_id = val("project_id")
        project = db.get(Project, int(project_id)) if project_id.isdigit() else None
        if project is None:
            project = db.query(Project).filter(Project.developer_id==developer.id, Project.name.ilike(project_name)).first()
        if project is None:
            project = Project(developer_id=developer.id, name=project_name)
            db.add(project)
        else:
            project.developer_id = developer.id
            project.name = project_name

        def setp(attr,name):
            x=val(name)
            if x: setattr(project,attr,x)
        setp("state","state"); setp("district","district"); setp("city","city")
        setp("locality","locality"); setp("address","address"); setp("pincode","pincode")
        setp("rera","rera"); setp("construction_status","tower_status")
        project.possession_date = _to_date(val("possession_date")) or project.possession_date
        project.brochure = val("brochure_url") or project.brochure
        amenities=[x.strip() for x in val("amenities").split(",") if x.strip()]
        if amenities: project.amenities_json=json.dumps(amenities,ensure_ascii=False)

        metadata={
            "v5": {
                "property_type":val("property_type"), "listing_status":val("listing_status"),
                "availability":val("availability"), "project_status":val("project_status"),
                "maps_link":val("maps_link"), "sub_locality":val("sub_locality"),
                "nearby_connectivity":val("nearby_connectivity"), "rera_url":val("rera_url"),
                "approved_banks":val("approved_banks"), "maintenance":val("maintenance"),
                "payment_plan":val("payment_plan"), "cp_commission":val("cp_commission"),
                "data_source":val("data_source"), "description":val("description"),
            }
        }
        project.description=json.dumps(metadata,ensure_ascii=False)
        db.flush()

        def lists(name): return vals(name)
        cb,ca,cr,ct,cp,cf,cv,ctx = [lists(x) for x in ("config_bhk","config_area","config_rate","config_total","config_parking","config_furnishing","config_availability","config_tower")]
        for i in range(max(len(cb),len(ca),len(cr),len(ct),len(cp),len(cf),len(cv),len(ctx))):
            bhk=cb[i] if i<len(cb) else ""
            if not bhk: continue
            area=_num(ca[i] if i<len(ca) else None)
            cfg=db.query(PropertyConfiguration).filter(PropertyConfiguration.project_id==project.id,PropertyConfiguration.configuration==bhk,PropertyConfiguration.carpet_area==area).first()
            if cfg is None:
                cfg=PropertyConfiguration(project_id=project.id,configuration=bhk,carpet_area=area)
                db.add(cfg)
            cfg.rate_per_sqft=_num(cr[i] if i<len(cr) else None)
            cfg.price_including_all=_num(ct[i] if i<len(ct) else None)
            cfg.availability=cv[i] if i<len(cv) else None
            cfg.source_details=json.dumps({"parking":cp[i] if i<len(cp) else "", "furnishing":cf[i] if i<len(cf) else "", "tower_wing":ctx[i] if i<len(ctx) else "", "source":val("data_source")},ensure_ascii=False)

        tn,tf,tp,ts=[lists(x) for x in ("tower_name","tower_floors","tower_possession","tower_status")]
        tower_by_name={}
        for i in range(max(len(tn),len(tf),len(tp),len(ts))):
            name=tn[i] if i<len(tn) else ""
            if not name: continue
            tower=db.query(Tower).filter(Tower.project_id==project.id,Tower.name.ilike(name)).first()
            if tower is None:
                tower=Tower(project_id=project.id,name=name); db.add(tower)
            if i<len(tf) and tf[i]:
                try: tower.floors=int(float(tf[i]))
                except ValueError: pass
            tower.possession_date=_to_date(tp[i]) if i<len(tp) and tp[i] else tower.possession_date
            tower.details=json.dumps({"construction_status":ts[i] if i<len(ts) else ""},ensure_ascii=False)
            db.flush(); tower_by_name[name.casefold()]=tower

        un,ut,uf,ub,ua,up,uk,ufurn,uav=[lists(x) for x in ("unit_number","unit_tower","unit_floor","unit_bhk","unit_area","unit_price","unit_parking","unit_furnishing","unit_availability")]
        for i in range(max(len(un),len(ut),len(uf),len(ub),len(ua),len(up),len(uk),len(ufurn),len(uav))):
            number=un[i] if i<len(un) else ""
            if not number: continue
            tower_name=ut[i] if i<len(ut) else ""
            tower=tower_by_name.get(tower_name.casefold()) if tower_name else None
            if tower is None and tower_name:
                tower=db.query(Tower).filter(Tower.project_id==project.id,Tower.name.ilike(tower_name)).first()
            unit=db.query(Unit).filter(Unit.project_id==project.id,Unit.unit_number==number).first()
            if unit is None:
                unit=Unit(project_id=project.id,unit_number=number); db.add(unit)
            if tower: unit.tower_id=tower.id
            if i<len(uf) and uf[i]:
                try: unit.floor=int(float(uf[i]))
                except ValueError: pass
            for attr,arr in (("bhk",ub),("carpet_area",ua),("price",up),("parking",uk),("furnishing",ufurn),("availability",uav)):
                if i<len(arr) and arr[i]:
                    setattr(unit,attr,_num(arr[i]) if attr in {"carpet_area","price"} else arr[i])
            unit.other_details=json.dumps({"source":val("data_source")},ensure_ascii=False)

        nn,nname,nd,ntime,ntransport=[lists(x) for x in ("nearby_category","nearby_name","nearby_distance","nearby_time","nearby_transport")]
        for i in range(max(len(nn),len(nname),len(nd),len(ntime),len(ntransport))):
            category=nn[i] if i<len(nn) else ""
            name=nname[i] if i<len(nname) else ""
            if not category or not name: continue
            nearby=db.query(NearbyLocation).filter(NearbyLocation.project_id==project.id,NearbyLocation.category.ilike(category),NearbyLocation.name.ilike(name)).first()
            if nearby is None:
                nearby=NearbyLocation(project_id=project.id,category=category,name=name); db.add(nearby)
            nearby.distance=nd[i] if i<len(nd) else None
            nearby.details=json.dumps({"travel_time":ntime[i] if i<len(ntime) else "", "transport_mode":ntransport[i] if i<len(ntransport) else ""},ensure_ascii=False)

        db.add(AuditEvent(actor_user_id=actor.id,action="PROPERTY_V5_MANUAL_SAVE",entity_type="Projects",entity_id=str(project.id),details=json.dumps({"project":project.name,"source":val("data_source")},ensure_ascii=False)))
        db.commit()
        return RedirectResponse(f"/admin/property-data?entity=Projects&message=V5+Property+saved+successfully+%28Project+%23{project.id}%29",status_code=303)
    except Exception as exc:
        db.rollback()
        return RedirectResponse("/admin/property-data/v5-entry?error="+quote_plus(str(exc)[:400]),status_code=303)


@router.post("/admin/property-data/load-godrej-nest", dependencies=[Depends(require_csrf)])
def load_godrej_nest_test_data(request: Request, db: Session = Depends(get_db)):
    """Explicitly load the reviewed Godrej Nest test inventory into Property Data."""
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    try:
        project_id, unit_count, gallery_count = seed_godrej_nest_test(db)
        db.add(AuditEvent(actor_user_id=actor.id, action="GODREJ_NEST_TEST_DATA_LOADED", entity_type="Projects", entity_id=str(project_id), details=json.dumps({"units": unit_count, "gallery_images": gallery_count, "source": "reviewed test seed"})))
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse("/admin/property-data?entity=Projects&error=Godrej+Nest+test+data+load+failed%3A+" + str(exc).replace(" ", "+")[:300], status_code=303)
    return RedirectResponse(f"/admin/property-data?entity=Projects&message=Godrej+Nest+test+property+loaded+%28project+%23{project_id}%2C+{unit_count}+units%2C+{gallery_count}+gallery+images%29", status_code=303)


@router.get("/admin/property-data/template")
def download_property_template(request: Request, db: Session = Depends(get_db)):
    if not _guard(request, db):
        return RedirectResponse("/", status_code=303)
    workbook = _property_workbook()
    return StreamingResponse(workbook, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=MetroCRM_Property_Data_Template.xlsx"})


@router.post("/admin/property-data/import", dependencies=[Depends(require_csrf)])
async def import_property_data(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return RedirectResponse("/admin/property-data?error=Only+.xlsx+files+are+supported", status_code=303)
    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
        if "QuickImport" in wb.sheetnames or any(name in wb.sheetnames for name in ("Developers","Projects","Configurations","Towers","Units","NearbyLocations")):
            created, updated, warnings = import_v5_workbook(
                db,
                (Developer, Project, PropertyConfiguration, Tower, Unit, NearbyLocation),
                wb,
                actor,
            )
            db.add(AuditEvent(
                actor_user_id=actor.id,
                action="PROPERTY_V5_WORKBOOK_IMPORTED",
                entity_type="PropertyData",
                entity_id=None,
                details=json.dumps({"created":created,"updated":updated,"warnings":warnings[:100]},ensure_ascii=False),
            ))
            db.commit()
            message=f"V5 import complete: {created} created, {updated} updated"
            if warnings:
                message += f", {len(warnings)} warnings"
            return RedirectResponse("/admin/property-data?entity=Projects&message="+quote_plus(message)[:800], status_code=303)
        if "Property Master" in wb.sheetnames:
            source_ws = wb["Property Master"]
            source_headers = [str(x.value).strip() if x.value is not None else "" for x in source_ws[1]]
            if "Project ID" in source_headers and "Configurations / Inventory" in source_headers:
                from app.models import Project
                existing_by_ref = {}
                for existing_project in db.scalars(select(Project)).all():
                    try:
                        desc = json.loads(existing_project.description or "{}")
                        ref = str((desc.get("property_master") or {}).get("property_ref") or "").strip()
                        if ref:
                            existing_by_ref[ref] = existing_project.id
                    except (TypeError, ValueError):
                        continue
                created = updated = 0
                for row_no, values in enumerate(source_ws.iter_rows(min_row=2, values_only=True), start=2):
                    raw = {source_headers[i]: values[i] for i in range(min(len(source_headers), len(values))) if source_headers[i]}
                    if not any(v not in (None, "") for v in raw.values()):
                        continue
                    ref = str(raw.get("Project ID") or "").strip()
                    if not ref:
                        raise ValueError(f"Property Master row {row_no}: Project ID is required")
                    raw_configs = json.loads(raw.get("Configurations / Inventory") or "[]")
                    if not isinstance(raw_configs, list):
                        raise ValueError(f"Property Master row {row_no}: Configurations / Inventory must be a JSON list")
                    configs = []
                    for config in raw_configs:
                        if isinstance(config, dict):
                            configs.extend(_source_configuration_rows(config, raw.get("Availability")))
                    tower_rows = _source_tower_rows(raw_configs)
                    rm_parts = []
                    for label in ("RM Mobile", "RM WhatsApp", "RM Email"):
                        value = str(raw.get(label) or "").strip()
                        if value:
                            rm_parts.append(f"{label}: {value}")
                    payload = {
                        "property_ref": ref, "developer": raw.get("Developer / Builder"),
                        "project_name": raw.get("Project Name"), "property_type": raw.get("Property Type"),
                        "availability_status": raw.get("Availability"), "full_address": raw.get("Full Address"),
                        "city": raw.get("City"), "locality": raw.get("Locality / Area"), "pincode": raw.get("Pincode"),
                        "google_maps_link": raw.get("Google Maps Link"), "nearby_connectivity": raw.get("Nearby Connectivity / Landmarks"),
                        "project_status": raw.get("Project Status"), "possession_date": raw.get("Possession Date"),
                        "total_towers": len(tower_rows) or None, "tower_details": tower_rows,
                        "rera_registration_number": raw.get("RERA Registration Number"), "approved_banks": raw.get("Approved Banks List"),
                        "title_deed_type": raw.get("Title Deed Type"), "configurations": configs,
                        "additional_government_charges": raw.get("Additional / Government Charges"), "payment_plan": raw.get("Payment Plan"),
                        "cp_commission_percent": raw.get("CP Commission %"), "site_flat_photos": raw.get("High-Res Site / Flat Photos"),
                        "floor_plan": raw.get("Floor Plan / Master Layout"), "master_layout": raw.get("Floor Plan / Master Layout"),
                        "project_pdf_brochure": raw.get("Project PDF Brochure"), "walkthrough_video_link": raw.get("Walkthrough Video Link"),
                        "rm_name": raw.get("RM Name"), "rm_contact": " | ".join(rm_parts) or None,
                        "site_visit_timings": raw.get("Site Visit Timings"), "booking_process": raw.get("Booking Process"),
                        "last_updated_date": raw.get("Last Updated Date"), "data_source": raw.get("Data Source"),
                    }
                    project_id = existing_by_ref.get(ref)
                    project, _ = apply_property_master(db, payload=payload, source=PropertyUpdateSource.BULK, actor=actor, project_id=project_id, create=project_id is None)
                    if project_id is None:
                        created += 1
                    else:
                        updated += 1
                    existing_by_ref[ref] = project.id
                db.commit()
                db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_DATA_IMPORTED", entity_type="PropertyMaster", entity_id=None, details=json.dumps({"filename": file.filename, "created": created, "updated": updated, "format": "metro_properties_32_field_master"})))
                db.commit()
                return RedirectResponse(f"/admin/property-data?message=Imported+{created}+new+and+updated+{updated}+Metro+Properties+records", status_code=303)
        wb = load_workbook(BytesIO(content), data_only=True)
        if "Property Master" not in wb.sheetnames:
            raise ValueError("Property Master sheet is required")
        master_ws = wb["Property Master"]
        cfg_ws = wb["Property Configurations"] if "Property Configurations" in wb.sheetnames else None

        master_headers = [str(x.value).strip() if x.value is not None else "" for x in master_ws[1]]
        cfg_headers = [str(x.value).strip() if x.value is not None else "" for x in cfg_ws[1]] if cfg_ws else []
        master_allowed = {"property_id"} | {field.key for field in PROPERTY_MASTER_FIELDS if field.key not in {"configurations", "tower_details"}}
        cfg_allowed = {"property_id", "property_ref", "configuration", "carpet_area", "price_including_all", "rate_per_sqft", "negotiation", "availability"}
        if any(h not in master_allowed for h in master_headers if h):
            raise ValueError("Property Master contains unsupported columns")
        if cfg_ws and any(h not in cfg_allowed for h in cfg_headers if h):
            raise ValueError("Property Configurations contains unsupported columns")

        config_by_ref = {}
        tower_by_ref = {}
        tower_ws = wb["Property Towers"] if "Property Towers" in wb.sheetnames else None
        tower_headers = [str(x.value).strip() if x.value is not None else "" for x in tower_ws[1]] if tower_ws else []
        for row_no, values in enumerate(tower_ws.iter_rows(min_row=2, values_only=True), start=2) if tower_ws else []:
            raw = {tower_headers[i]: values[i] for i in range(min(len(tower_headers), len(values))) if tower_headers[i]}
            if not any(v not in (None, "") for v in raw.values()): continue
            keys = [str(raw.get(x) or "").strip() for x in ("property_id","property_ref") if str(raw.get(x) or "").strip()]
            if not keys: raise ValueError(f"Property Towers row {row_no}: property_id or property_ref is required")
            tower_row = {"tower_name": raw.get("tower_name"), "total_floors": raw.get("total_floors"), "possession_date": raw.get("possession_date")}
            for key in keys: tower_by_ref.setdefault(key, []).append(dict(tower_row))

        for row_no, values in enumerate(cfg_ws.iter_rows(min_row=2, values_only=True), start=2) if cfg_ws else []:
            raw = {cfg_headers[i]: values[i] for i in range(min(len(cfg_headers), len(values))) if cfg_headers[i]}
            if not any(v not in (None, "") for v in raw.values()):
                continue
            ref = str(raw.get("property_ref") or "").strip()
            pid = str(raw.get("property_id") or "").strip()
            keys = [x for x in (pid, ref) if x]
            if not keys:
                raise ValueError(f"Property Configurations row {row_no}: property_id or property_ref is required")
            config_row = {
                "configuration": raw.get("configuration"),
                "carpet_area": raw.get("carpet_area"),
                "price_including_all": raw.get("price_including_all"),
                "rate_per_sqft": raw.get("rate_per_sqft"),
                "negotiation": raw.get("negotiation"),
                "availability": raw.get("availability"),
            }
            for key in keys:
                config_by_ref.setdefault(key, []).append(dict(config_row))

        created = updated = 0
        refs_to_ids = {}
        for row_no, values in enumerate(master_ws.iter_rows(min_row=2, values_only=True), start=2):
            raw = {master_headers[i]: values[i] for i in range(min(len(master_headers), len(values))) if master_headers[i]}
            if not any(v not in (None, "") for v in raw.values()):
                continue
            property_id = str(raw.get("property_id") or "").strip()
            property_ref = str(raw.get("property_ref") or "").strip()
            if not property_id and not property_ref:
                raise ValueError(f"Property Master row {row_no}: property_id or property_ref is required")
            key = property_id or property_ref
            payload = {k: v for k, v in raw.items() if k not in {"property_id", "property_ref"}}
            payload["developer"] = raw.get("developer")
            payload["configurations"] = config_by_ref.get(key, [])
            payload["tower_details"] = tower_by_ref.get(key, [])
            project, _ = apply_property_master(
                db, payload=payload, source=PropertyUpdateSource.BULK, actor=actor,
                project_id=int(property_id) if property_id else None, create=not bool(property_id)
            )
            refs_to_ids[property_ref] = project.id if property_ref else project.id
            if property_id:
                updated += 1
            else:
                created += 1
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse("/admin/property-data?error=" + str(exc).replace(" ", "+")[:350], status_code=303)
    db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_DATA_IMPORTED", entity_type="PropertyMaster", entity_id=None, details=json.dumps({"filename": file.filename, "created": created, "updated": updated, "format": "canonical_property_master_v3"})))
    db.commit()
    return RedirectResponse(f"/admin/property-data?message=Imported+{created}+new+and+updated+{updated}+properties+using+Property+Master+format", status_code=303)


@router.post("/admin/property-data/bulk-modify", dependencies=[Depends(require_csrf)])
async def bulk_modify_property_data(request: Request, db: Session = Depends(get_db)):
    """Apply explicitly supplied fields to multiple selected Projects."""
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    form = await request.form()
    raw_ids = form.getlist("project_ids")
    project_ids = []
    for raw_id in raw_ids:
        try:
            project_ids.append(int(str(raw_id)))
        except (TypeError, ValueError):
            continue
    project_ids = list(dict.fromkeys(project_ids))
    if not project_ids:
        return RedirectResponse("/admin/property-data?entity=Projects&error=Select+at+least+one+property", status_code=303)

    allowed_fields = ("city", "locality", "address", "rera")
    values = {field: str(form.get(field, "")).strip() for field in allowed_fields}
    supplied = {field: value for field, value in values.items() if value}
    if not supplied:
        return RedirectResponse("/admin/property-data?entity=Projects&error=Enter+at+least+one+field+to+modify", status_code=303)

    try:
        rows = db.scalars(select(PROPERTY_MODELS["Projects"][0]).where(PROPERTY_MODELS["Projects"][0].id.in_(project_ids))).all()
        found = {row.id: row for row in rows}
        missing = [str(pid) for pid in project_ids if pid not in found]
        if missing:
            raise ValueError("Property ID not found: " + ", ".join(missing[:10]))
        for row in rows:
            apply_property_master(
                db,
                payload=supplied,
                source=PropertyUpdateSource.BULK,
                actor=actor,
                project_id=row.id,
                create=False,
            )
            for field, value in supplied.items():
                if field in {"city", "locality", "address", "rera"}:
                    setattr(row, field, value)
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse("/admin/property-data?entity=Projects&error=" + str(exc).replace(" ", "+")[:350], status_code=303)
    return RedirectResponse("/admin/property-data?entity=Projects&message=Updated+" + str(len(rows)) + "+properties+successfully", status_code=303)


@router.post("/super-admin/property-data/projects/{project_id}/delete", dependencies=[Depends(require_csrf)])
def delete_single_property_data(project_id: int, request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor or actor.role != "SUPER_ADMIN":
        return RedirectResponse("/dashboard", status_code=303)
    if not db.get(Project, project_id):
        return RedirectResponse("/super-admin/property-data?entity=Projects&error=Property+not+found", status_code=303)
    return _delete_single_property_data(project_id, actor, db)


def _delete_single_property_data(project_id: int, actor: User, db: Session):
    try:
        if db.scalar(select(Booking.id).where(Booking.project_id == project_id).limit(1)) or db.scalar(select(Commission.id).where(Commission.project_id == project_id).limit(1)) or db.scalar(select(SiteVisit.id).where(SiteVisit.project_id == project_id).limit(1)) or db.scalar(select(Token.id).where(Token.project_id == project_id).limit(1)):
            raise ValueError("Cannot delete property with linked booking, commission, site visit or token records")
        unit_ids = db.scalars(select(Unit.id).where(Unit.project_id == project_id)).all()
        db.query(NearbyLocation).filter(NearbyLocation.project_id == project_id).delete(synchronize_session=False)
        if unit_ids:
            db.query(Token).filter(Token.unit_id.in_(unit_ids)).update({Token.unit_id: None}, synchronize_session=False)
            db.query(Booking).filter(Booking.unit_id.in_(unit_ids)).update({Booking.unit_id: None}, synchronize_session=False)
        db.query(Unit).filter(Unit.project_id == project_id).delete(synchronize_session=False)
        db.query(Tower).filter(Tower.project_id == project_id).delete(synchronize_session=False)
        db.query(PropertyConfiguration).filter(PropertyConfiguration.project_id == project_id).delete(synchronize_session=False)
        db.query(Project).filter(Project.id == project_id).delete(synchronize_session=False)
        db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_RECORD_DELETED", entity_type="Projects", entity_id=str(project_id), details=json.dumps({"property_only_records": True})))
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse("/super-admin/property-data?entity=Projects&error=" + str(exc).replace(" ", "+")[:350], status_code=303)
    return RedirectResponse("/super-admin/property-data?entity=Projects&message=Property+" + str(project_id) + "+deleted+successfully", status_code=303)


@router.post("/admin/property-data/save", dependencies=[Depends(require_csrf)])
async def save_property_data(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    form = await request.form()
    entity = str(form.get("entity", "Projects"))
    if entity not in PROPERTY_MODELS:
        return RedirectResponse("/admin/property-data?error=Invalid+property+type", status_code=303)
    model, fields = PROPERTY_MODELS[entity]
    raw_id = str(form.get("id", "")).strip()
    try:
        obj = db.get(model, int(raw_id)) if raw_id else model()
        if raw_id and obj is None:
            raise ValueError("Record ID does not exist")
        if entity == "Projects":
            canonical_payload = {k: v for k, v in form.items() if k not in {"csrf_token", "entity", "id"}}
            if any(k in canonical_payload for k in ("project_name", "property_type", "configuration", "carpet_area", "possession_date", "full_address", "bsp", "pincode", "rera_registration_number")):
                obj, _ = apply_property_master(
                    db,
                    payload=canonical_payload,
                    source=PropertyUpdateSource.INDIVIDUAL,
                    actor=actor,
                    project_id=int(raw_id) if raw_id else None,
                    create=not bool(raw_id),
                )
                db.commit()
                return RedirectResponse(f"/admin/property-data?entity=Projects&id={obj.id}&message=Property+Master+saved", status_code=303)
        columns = {column.name: column for column in model.__table__.columns}
        for field in fields:
            if field == "id" or field not in columns:
                continue
            value = form.get(field)
            converted = _convert_excel_value(columns[field], value)
            if field.endswith("_json") and converted is not None:
                json.loads(converted)
            setattr(obj, field, converted)
        db.add(obj)
        db.flush()
        action = "PROPERTY_RECORD_UPDATED" if raw_id else "PROPERTY_RECORD_CREATED"
        db.add(AuditEvent(actor_user_id=actor.id, action=action, entity_type=entity, entity_id=str(obj.id), details=json.dumps({"online_editor": True})))
        db.commit()
    except Exception as exc:
        db.rollback()
        return RedirectResponse(f"/admin/property-data?entity={entity}&error=" + str(exc).replace(" ", "+")[:350], status_code=303)
    return RedirectResponse(f"/admin/property-data?entity={entity}&id={obj.id}&message=Property+record+saved", status_code=303)


def _create_export_request(actor: User, db: Session):
    entities = list(PROPERTY_SHEETS)
    pending = db.scalar(select(DataRequest).where(DataRequest.requester_id == actor.id, DataRequest.action == "EXPORT", DataRequest.entity_type == "PropertyData", DataRequest.status == "PENDING").order_by(DataRequest.id.desc()))
    if pending:
        return False
    request_row = DataRequest(requester_id=actor.id, action="EXPORT", entity_type="PropertyData", filters_json=json.dumps({"entities": entities}), reason="Admin requested property data export", status="PENDING")
    db.add(request_row)
    db.flush()
    db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_EXPORT_REQUESTED", entity_type="PropertyData", entity_id=str(request_row.id), details=json.dumps({"entities": entities, "status": "PENDING"})))
    db.commit()
    return True


@router.get("/admin/property-data/export")
def export_property_data(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    if actor.role == "ADMIN":
        return RedirectResponse("/admin/property-data?error=Admin+must+submit+an+export+approval+request", status_code=303)
    entities = list(PROPERTY_SHEETS)
    workbook = _export_workbook(entities, db)
    db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_DATA_EXPORTED", entity_type="PropertyData", entity_id=None, details=json.dumps({"entities": entities, "approval": "direct"})))
    db.commit()
    return StreamingResponse(workbook, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=MetroCRM_Property_Data_Export.xlsx"})


@router.post("/admin/property-data/export/request", dependencies=[Depends(require_csrf)])
def request_property_export(request: Request, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    if actor.role != "ADMIN":
        return RedirectResponse("/admin/property-data/export", status_code=303)
    created = _create_export_request(actor, db)
    if not created:
        return RedirectResponse("/admin/property-data?error=Export+approval+already+pending", status_code=303)
    return RedirectResponse("/admin/property-data?message=Export+request+sent+to+Super+Admin+for+approval", status_code=303)


@router.get("/admin/property-data/export/download/{request_id}")
def download_approved_property_export(request: Request, request_id: int, db: Session = Depends(get_db)):
    actor = _guard(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    row = db.get(DataRequest, request_id)
    if not row or row.action != "EXPORT" or row.entity_type != "PropertyData" or row.status != "APPROVED":
        return RedirectResponse("/admin/property-data?error=Export+approval+not+available", status_code=403)
    if actor.role != "SUPER_ADMIN" and row.requester_id != actor.id:
        return RedirectResponse("/admin/property-data?error=You+can+only+download+your+own+approved+export", status_code=403)
    try:
        details = json.loads(row.filters_json or "{}")
        entities = [x for x in details.get("entities", []) if x in PROPERTY_SHEETS]
    except (TypeError, ValueError):
        entities = list(PROPERTY_SHEETS)
    workbook = _export_workbook(entities or list(PROPERTY_SHEETS), db)
    db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_DATA_EXPORTED", entity_type="PropertyData", entity_id=str(row.id), details=json.dumps({"entities": entities, "approval_request_id": row.id, "approved_by": row.decided_by_id})))
    db.commit()
    return StreamingResponse(workbook, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=MetroCRM_Property_Data_Export.xlsx"})


@router.get("/super-admin/property-data-export-requests", response_class=HTMLResponse)
def property_export_requests(request: Request, db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor or actor.role != "SUPER_ADMIN":
        return RedirectResponse("/dashboard", status_code=303)
    rows = db.scalars(select(DataRequest).where(DataRequest.action == "EXPORT", DataRequest.entity_type == "PropertyData").order_by(DataRequest.id.desc()).limit(100)).all()
    users = {u.id: u for u in db.scalars(select(User).where(User.id.in_({r.requester_id for r in rows}))).all()} if rows else {}
    html = ["<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>MetroCRM | Property Export Approvals</title><style>body{font-family:Inter,Arial,sans-serif;background:#f5f7fb;margin:0;padding:32px;color:#10213f}.wrap{max-width:1050px;margin:auto}.card{background:#fff;border-radius:16px;padding:24px;box-shadow:0 8px 30px #0001}.row{padding:16px 0;border-top:1px solid #e5eaf2}.actions{display:flex;gap:8px;flex-wrap:wrap}.btn{border:0;border-radius:9px;padding:9px 14px;font-weight:700;cursor:pointer}.approve{background:#16803c;color:#fff}.reject{background:#c62828;color:#fff}.muted{color:#64748b;font-size:12px}.pill{display:inline-block;background:#eff6ff;color:#1d4ed8;border-radius:999px;padding:4px 9px;font-size:11px;font-weight:700}</style></head><body><div class='wrap'><div class='card'><h1>Property Data Export Approvals</h1><p class='muted'>Admin exports require explicit Super Admin approval. Approved requests can then be downloaded by the requesting Admin.</p>"]
    for row in rows:
        user = users.get(row.requester_id)
        status = row.status
        if status == "PENDING":
            actions = f"<div class='actions'><form method='post' action='/super-admin/property-data-export-requests/{row.id}/approve'><input type='hidden' name='csrf_token' value='{csrf_token(request)}'><button class='btn approve'>Approve Export</button></form><form method='post' action='/super-admin/property-data-export-requests/{row.id}/reject'><input type='hidden' name='csrf_token' value='{csrf_token(request)}'><button class='btn reject'>Reject</button></form></div>"
        else:
            actions = ""
        html.append(f"<div class='row'><b>{user.username if user else 'Unknown user'}</b> · <span class='pill'>{status}</span><div class='muted'>Request #{row.id} · {row.created_at} · {row.reason}</div>{actions}</div>")
    if not rows:
        html.append("<div class='row muted'>No property export requests.</div>")
    html.append("</div></div></body></html>")
    return HTMLResponse("".join(html))


def _decide_property_export(actor: User | None, request_id: int, approve: bool, db: Session):
    if not actor or actor.role != "SUPER_ADMIN":
        return RedirectResponse("/dashboard", status_code=303)
    row = db.get(DataRequest, request_id)
    if not row or row.action != "EXPORT" or row.entity_type != "PropertyData" or row.status != "PENDING":
        return RedirectResponse("/super-admin/property-data-export-requests", status_code=303)
    row.status = "APPROVED" if approve else "REJECTED"
    row.decided_by_id = actor.id
    row.decided_at = datetime.now(timezone.utc).replace(tzinfo=None)
    row.approval_reason = "Approved by Super Admin" if approve else "Rejected by Super Admin"
    db.add(AuditEvent(actor_user_id=actor.id, action="PROPERTY_EXPORT_APPROVED" if approve else "PROPERTY_EXPORT_REJECTED", entity_type="PropertyData", entity_id=str(row.id), details=json.dumps({"requester_id": row.requester_id})))
    db.commit()
    return RedirectResponse("/super-admin/property-data-export-requests", status_code=303)


@router.post("/super-admin/property-data-export-requests/{request_id}/approve", dependencies=[Depends(require_csrf)])
def approve_property_export(request: Request, request_id: int, db: Session = Depends(get_db)):
    return _decide_property_export(_user(request, db), request_id, True, db)


@router.post("/super-admin/property-data-export-requests/{request_id}/reject", dependencies=[Depends(require_csrf)])
def reject_property_export(request: Request, request_id: int, db: Session = Depends(get_db)):
    return _decide_property_export(_user(request, db), request_id, False, db)
