from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from html import escape
from io import BytesIO
import json

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.csrf import csrf_token, require_csrf
from app.database import get_db
from app.models import AuditEvent, Booking, Commission, FollowUp, Lead, SiteVisit, Token, User

router = APIRouter(tags=["advanced-reports"])

STAFF_REPORT_ROLES = {"TL", "MANAGER"}
ADMIN_ROLES = {"ADMIN", "SUPER_ADMIN"}
PERIODS = {"daily", "weekly", "monthly"}
BUSINESS_PERIODS = {"weekly", "monthly", "quarterly", "yearly", "custom"}


def _user(request: Request, db: Session) -> User | None:
    from app.main import _current_user
    return _current_user(request, db)


def _bounds(period: str, from_date: str | None = None, to_date: str | None = None) -> tuple[datetime, datetime, str]:
    today = date.today()
    if period == "daily":
        start = end = today
    elif period == "weekly":
        start = today - timedelta(days=today.weekday())
        end = start + timedelta(days=6)
    elif period == "monthly":
        start = today.replace(day=1)
        next_month = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
        end = next_month - timedelta(days=1)
    elif period == "quarterly":
        month = ((today.month - 1) // 3) * 3 + 1
        start = date(today.year, month, 1)
        end = (date(today.year + (1 if month == 10 else 0), 1 if month == 10 else month + 3, 1) - timedelta(days=1))
    elif period == "yearly":
        start = date(today.year, 1, 1)
        end = date(today.year, 12, 31)
    elif period == "custom":
        try:
            start = date.fromisoformat(from_date or "")
            end = date.fromisoformat(to_date or "")
        except ValueError:
            raise ValueError("Invalid custom date range")
        if end < start:
            raise ValueError("To date cannot be before From date")
    else:
        raise ValueError("Unsupported report period")
    return datetime.combine(start, datetime.min.time()), datetime.combine(end, datetime.max.time()), f"{start.isoformat()} to {end.isoformat()}"


def _scope_ids(actor: User, db: Session) -> set[int]:
    if actor.role in ADMIN_ROLES:
        return set(db.scalars(select(User.id)).all())
    ids = {actor.id}
    changed = True
    while changed:
        changed = False
        children = db.scalars(select(User).where(User.reporting_manager_id.in_(ids))).all()
        for child in children:
            if child.id not in ids:
                ids.add(child.id)
                changed = True
    return ids


def _staff_report(actor: User, db: Session, start: datetime, end: datetime) -> dict:
    ids = _scope_ids(actor, db)
    leads = db.scalars(select(Lead).where(Lead.assigned_user_id.in_(ids), Lead.created_at.between(start, end))).all()
    followups = db.scalars(select(FollowUp).where(FollowUp.assigned_user_id.in_(ids), FollowUp.due_at.between(start, end))).all()
    visits = db.scalars(select(SiteVisit).where(SiteVisit.scheduled_at.between(start, end), or_(SiteVisit.sales_employee_id.in_(ids), SiteVisit.conducted_by_id.in_(ids)))).all()
    bookings = db.scalars(select(Booking).where(Booking.booking_date.between(start, end), or_(Booking.employee_id.in_(ids), Booking.manager_id.in_(ids)))).all()
    tokens = db.scalars(select(Token).where(Token.token_date.between(start, end), or_(Token.salesperson_id.in_(ids), Token.closing_manager_id.in_(ids)))).all()
    commissions = db.scalars(select(Commission).where(Commission.payment_date.between(start, end), Commission.employee_id.in_(ids))).all()
    return {
        "scope_ids": ids,
        "metrics": {
            "leads": len(leads), "new_leads": sum(x.status == "NEW" for x in leads),
            "closed_leads": sum(str(x.status).upper() in {"CLOSED", "WON", "CONVERTED"} for x in leads),
            "followups": len(followups), "completed_followups": sum(x.completed for x in followups),
            "site_visits": len(visits), "bookings": len(bookings),
            "booking_value": sum((x.sale_value or 0) for x in bookings),
            "token_value": sum((x.token_amount or 0) for x in tokens),
            "commission_received": sum((x.received_amount or 0) for x in commissions),
            "commission_pending": sum((x.pending_amount or 0) for x in commissions),
        },
    }


def _business_report(db: Session, start: datetime, end: datetime) -> dict:
    leads = db.scalar(select(func.count(Lead.id)).where(Lead.created_at.between(start, end))) or 0
    won = db.scalar(select(func.count(Lead.id)).where(Lead.created_at.between(start, end), Lead.status.in_(["CLOSED", "WON", "CONVERTED"]))) or 0
    bookings = db.scalars(select(Booking).where(Booking.booking_date.between(start, end))).all()
    tokens = db.scalars(select(Token).where(Token.token_date.between(start, end))).all()
    commissions = db.scalars(select(Commission).where(Commission.payment_date.between(start, end))).all()
    visits = db.scalar(select(func.count(SiteVisit.id)).where(SiteVisit.scheduled_at.between(start, end))) or 0
    followups = db.scalar(select(func.count(FollowUp.id)).where(FollowUp.due_at.between(start, end))) or 0
    return {
        "metrics": {
            "leads": leads, "won_leads": won,
            "conversion_rate": (won / leads * 100) if leads else 0,
            "site_visits": visits, "followups": followups,
            "bookings": len(bookings),
            "booking_amount": sum((x.booking_amount or 0) for x in bookings),
            "sales_value": sum((x.sale_value or 0) for x in bookings),
            "token_value": sum((x.token_amount or 0) for x in tokens),
            "refund_value": sum((x.refund_amount or 0) for x in tokens),
            "commission_amount": sum((x.amount or 0) for x in commissions),
            "commission_received": sum((x.received_amount or 0) for x in commissions),
            "commission_pending": sum((x.pending_amount or 0) for x in commissions),
        },
        "bookings": bookings, "tokens": tokens, "commissions": commissions,
    }


def _page(title: str, body: str) -> HTMLResponse:
    return HTMLResponse(f"<!doctype html><html><head><meta name='viewport' content='width=device-width,initial-scale=1'><title>MetroCRM | {escape(title)}</title><style>body{{font-family:Arial,sans-serif;background:#f5f7fb;color:#10213f;padding:24px}}.wrap{{max-width:1200px;margin:auto}}.card{{background:white;border:1px solid #e5e7eb;border-radius:14px;padding:20px;margin:0 0 18px}}.grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}}.stat{{border:1px solid #e5e7eb;border-radius:10px;padding:14px}}.v{{font-size:22px;font-weight:800;margin-top:5px}}.muted{{color:#64748b;font-size:13px}}input,select,button{{padding:9px;border:1px solid #cbd5e1;border-radius:7px}}button,.btn{{background:#2563eb;color:white;border:0;padding:10px 14px;border-radius:7px;text-decoration:none;cursor:pointer}}.row{{display:flex;gap:10px;flex-wrap:wrap;align-items:end}}label{{display:block;font-size:12px;font-weight:700;margin-bottom:5px}}@media(max-width:800px){{.grid{{grid-template-columns:repeat(2,1fr)}}}}</style></head><body><div class='wrap'><h1>{escape(title)}</h1>{body}<p><a class='btn' href='/dashboard'>Dashboard</a></p></div></body></html>")


def _metric_html(metrics: dict) -> str:
    return "<div class='grid'>" + "".join(f"<div class='stat'><div class='muted'>{escape(k.replace('_',' ').title())}</div><div class='v'>{escape(str(float(v)) if isinstance(v, Decimal) else str(round(v,2)) if isinstance(v,float) else str(v))}</div></div>" for k,v in metrics.items()) + "</div>"


def _xlsx(title: str, period_label: str, metrics: dict, rows: list[tuple[str,str,str]]) -> BytesIO:
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    ws.append([title]); ws.append(["Period", period_label]); ws.append([]); ws.append(["Metric", "Value"])
    for c in ws[1] + ws[4]: c.font = Font(bold=True)
    for k,v in metrics.items(): ws.append([k.replace('_',' ').title(), float(v) if isinstance(v, (Decimal,float)) else v])
    if rows:
        detail = wb.create_sheet("Transactions"); detail.append(["Type","Date","Amount"])
        for c in detail[1]: c.font = Font(bold=True)
        for r in rows: detail.append(list(r))
    out=BytesIO(); wb.save(out); out.seek(0); return out


@router.get("/reports/advanced", response_class=HTMLResponse)
def advanced_reports(request: Request, db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor: return RedirectResponse("/", status_code=303)
    if actor.role in ADMIN_ROLES: return RedirectResponse("/admin/business-reports", status_code=303)
    if actor.role not in STAFF_REPORT_ROLES: return RedirectResponse("/reports", status_code=303)
    csrf = csrf_token(request)
    body = "<div class='card'><h2>My Team Report</h2><p class='muted'>Daily, weekly or monthly report. Sharing is shown only as <b>Admin</b>; Super Admin access remains hidden from staff.</p><form class='row' method='post' action='/reports/advanced/share'><input type='hidden' name='csrf_token' value='"+csrf+"'><div><label>Period</label><select name='period'><option>daily</option><option>weekly</option><option>monthly</option></select></div><div><button type='submit'>Share with Admin</button></div></form></div>"
    return _page("Staff Reports", body)


@router.post("/reports/advanced/share", dependencies=[Depends(require_csrf)])
def share_staff_report(request: Request, period: str = Form(...), db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor or actor.role not in STAFF_REPORT_ROLES or period not in PERIODS:
        return RedirectResponse("/reports/advanced", status_code=303)
    start, end, label = _bounds(period)
    data = _staff_report(actor, db, start, end)
    details = {"period": period, "from": start.date().isoformat(), "to": end.date().isoformat(), "metrics": data["metrics"], "shared_to": "ADMIN", "scope": "actor_reporting_tree"}
    db.add(AuditEvent(actor_user_id=actor.id, action="STAFF_REPORT_SHARED", entity_type="Report", entity_id=actor.id.__str__(), details=json.dumps(details, default=str)))
    db.commit()
    return RedirectResponse("/reports/advanced?shared=1", status_code=303)


@router.get("/admin/staff-reports", response_class=HTMLResponse)
def admin_staff_reports(request: Request, db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor or actor.role not in ADMIN_ROLES: return RedirectResponse("/dashboard", status_code=303)
    rows = db.scalars(select(AuditEvent).where(AuditEvent.action == "STAFF_REPORT_SHARED").order_by(AuditEvent.id.desc()).limit(100)).all()
    body = "<div class='card'><p class='muted'>Staff reports shared to Admin. Super Admin receives the same reports through the secure executive view.</p><table><tr><th>Staff</th><th>Period</th><th>From</th><th>To</th><th>Leads</th><th>Bookings</th><th>Sales</th></tr>"
    for row in rows:
        d=json.loads(row.details or "{}")
        m=d.get("metrics",{})
        body += f"<tr><td>{row.actor_user_id}</td><td>{escape(str(d.get('period','')))}</td><td>{escape(str(d.get('from','')))}</td><td>{escape(str(d.get('to','')))}</td><td>{m.get('leads',0)}</td><td>{m.get('bookings',0)}</td><td>₹ {float(m.get('booking_value',0)):,.2f}</td></tr>"
    body += "</table></div>"
    return _page("Staff Reports Received", body)


@router.get("/admin/business-reports", response_class=HTMLResponse)
def business_reports(request: Request, period: str = "monthly", from_date: str | None = None, to_date: str | None = None, db: Session = Depends(get_db)):
    actor = _user(request, db)
    if not actor or actor.role not in ADMIN_ROLES: return RedirectResponse("/dashboard", status_code=303)
    try: start,end,label=_bounds(period,from_date,to_date)
    except ValueError as exc: return _page("Business Reports", f"<div class='card'>Invalid period: {escape(str(exc))}</div>")
    data=_business_report(db,start,end); csrf=csrf_token(request)
    body="<div class='card'><form class='row' method='get'><div><label>Period</label><select name='period'>"+"".join(f"<option value='{p}' {'selected' if p==period else ''}>{p.title()}</option>" for p in sorted(BUSINESS_PERIODS))+"</select></div><div><label>From date (Custom)</label><input type='date' name='from_date' value='"+escape(from_date or '')+"'></div><div><label>To date (Custom)</label><input type='date' name='to_date' value='"+escape(to_date or '')+"'></div><div><button type='submit'>Generate</button></div></form><p class='muted'>Period: "+escape(label)+"</p></div>"
    body += "<div class='card'>"+_metric_html(data['metrics'])+"</div>"
    body += f"<div class='card'><form method='post' action='/admin/business-reports/export'><input type='hidden' name='csrf_token' value='{csrf}'><input type='hidden' name='period' value='{escape(period)}'><input type='hidden' name='from_date' value='{escape(from_date or '')}'><input type='hidden' name='to_date' value='{escape(to_date or '')}'><button type='submit'>Download Excel Report</button></form><p class='muted'>Includes management, sales, booking, token/refund and commission fields useful for accounting/ITR/GST preparation. Tax filing should be reconciled with statutory books and your tax professional.</p></div>"
    return _page("Business Reports", body)


@router.post("/admin/business-reports/export", dependencies=[Depends(require_csrf)])
def export_business_report(request: Request, period: str = Form(...), from_date: str = Form(""), to_date: str = Form(""), db: Session = Depends(get_db)):
    actor=_user(request,db)
    if not actor or actor.role not in ADMIN_ROLES: return RedirectResponse("/dashboard",status_code=303)
    try: start,end,label=_bounds(period,from_date or None,to_date or None)
    except ValueError: return RedirectResponse("/admin/business-reports",status_code=303)
    data=_business_report(db,start,end)
    rows=[]
    for x in data['bookings']: rows.append(("BOOKING",x.booking_date.isoformat() if x.booking_date else "",str(x.sale_value or 0)))
    for x in data['tokens']: rows.append(("TOKEN",x.token_date.isoformat() if x.token_date else "",str(x.token_amount or 0)))
    for x in data['commissions']: rows.append(("COMMISSION",x.payment_date.isoformat() if x.payment_date else "",str(x.amount or 0)))
    wb=_xlsx("MetroCRM Business Report",label,data['metrics'],rows)
    return StreamingResponse(wb,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=MetroCRM_Business_Report.xlsx"})


@router.get("/super-admin/business-reports", response_class=HTMLResponse)
def super_admin_business_reports(request: Request, period: str = "monthly", from_date: str | None = None, to_date: str | None = None, db: Session = Depends(get_db)):
    actor=_user(request,db)
    if not actor or actor.role != "SUPER_ADMIN": return RedirectResponse("/super-admin/login",status_code=303)
    start,end,label=_bounds(period,from_date,to_date); data=_business_report(db,start,end)
    rows=db.scalars(select(AuditEvent).where(AuditEvent.action=="STAFF_REPORT_SHARED").order_by(AuditEvent.id.desc()).limit(100)).all()
    body="<div class='card'><p class='muted'>Executive view. Staff users and Managers are never shown this Super Admin destination.</p>"+_metric_html(data['metrics'])+"</div>"
    body += "<div class='card'><h2>Staff reports received</h2><p>Shared staff reports are retained for executive review.</p><p><a class='btn' href='/admin/staff-reports'>Open staff report register</a></p><p>Total shared reports: "+str(len(rows))+"</p></div>"
    return _page("Executive Business Reports",body)
