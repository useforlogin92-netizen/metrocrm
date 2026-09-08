from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
import json
from html import escape
from zipfile import BadZipFile, ZipFile, is_zipfile
from typing import Any

from fastapi import APIRouter, Depends, Form, Request, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from app.csrf import csrf_token, require_csrf
from app.database import get_db

router = APIRouter(tags=["data-tools"])

IMPORT_SHEETS: dict[str, tuple[type, list[str], dict[str, str]]] = {}


def _models():
    from app.models import (
        Customer, CustomerRequirement, Lead, Developer, Project, Tower, Unit,
        FollowUp, SiteVisit, Token, Booking, Commission, DataRequest, AuditEvent,
    )
    return {
        "Customers": (Customer, ["id", "name", "mobile", "alternate_mobile", "email", "notes"]),
        "CustomerRequirements": (CustomerRequirement, ["id", "customer_id", "buy_rent", "property_category", "locations_json", "budget_min", "budget_max", "bhk", "property_type", "carpet_area_min", "carpet_area_max", "possession_requirement", "furnishing", "parking", "project_preference", "developer_preference", "loan_required", "purpose", "notes", "is_active"]),
        "Leads": (Lead, ["id", "customer_id", "source", "status", "assigned_user_id", "original_owner_id", "requirement", "budget_min", "budget_max", "preferred_location", "bhk", "property_type", "buy_rent"]),
        "Developers": (Developer, ["id", "name", "contact", "details"]),
        "Projects": (Project, ["id", "developer_id", "name", "state", "district", "city", "locality", "sector", "address", "pincode", "latitude", "longitude", "rera", "possession_date", "builder_commitment_date", "construction_status", "amenities_json", "photos_json", "brochure", "description"]),
        "Towers": (Tower, ["id", "project_id", "name", "floors", "possession_date", "details"]),
        "Units": (Unit, ["id", "project_id", "tower_id", "unit_number", "floor", "total_building_floors", "bhk", "property_type", "carpet_area", "built_up_area", "super_built_up_area", "plot_area", "price", "price_per_sqft", "parking", "facing", "furnishing", "possession", "rera", "availability", "status", "latitude", "longitude", "amenities_json", "photos_json", "floor_plan", "brochure", "description", "other_details"]),
        "NearbyLocations": (__import__("app.models", fromlist=["NearbyLocation"]).NearbyLocation, ["id", "project_id", "category", "name", "distance", "details", "verified", "verified_at"]),
    }


def _report_models():
    m = _models()
    from app.models import FollowUp, SiteVisit, Token, Booking, Commission, User
    return m, FollowUp, SiteVisit, Token, Booking, Commission, User


def _current_user(request: Request, db: Session):
    from app.main import _current_user as get_current_user
    return get_current_user(request, db)


def _audit(db: Session, actor_id: int | None, action: str, entity_type: str | None = None, entity_id: str | None = None, details: dict | None = None):
    from app.models import AuditEvent
    db.add(AuditEvent(actor_user_id=actor_id, action=action, entity_type=entity_type, entity_id=entity_id, details=json.dumps(details or {})))


def _guard_admin(request: Request, db: Session):
    user = _current_user(request, db)
    if not user or user.role not in {"SUPER_ADMIN", "ADMIN"}:
        return None
    return user


def _guard_super_admin(request: Request, db: Session):
    user = _current_user(request, db)
    if not user or user.role != "SUPER_ADMIN":
        return None
    return user


def _template_headers(model, fields: list[str]) -> list[str]:
    return fields


def _sample_value(sheet: str, field: str):
    examples = {
        ("Customers", "name"): "Example Customer",
        ("Customers", "mobile"): "9876543210",
        ("Customers", "email"): "customer@example.com",
        ("CustomerRequirements", "customer_id"): 1,
        ("CustomerRequirements", "buy_rent"): "BUY",
        ("CustomerRequirements", "property_category"): "RESIDENTIAL",
        ("CustomerRequirements", "budget_min"): 5000000,
        ("CustomerRequirements", "budget_max"): 10000000,
        ("Leads", "source"): "WEBSITE",
        ("Leads", "status"): "NEW",
        ("Leads", "assigned_user_id"): 1,
        ("Leads", "preferred_location"): "Thane West",
        ("Leads", "budget_min"): 5000000,
        ("Leads", "budget_max"): 10000000,
        ("Developers", "name"): "Example Developer",
        ("Projects", "developer_id"): 1,
        ("Projects", "name"): "Example Heights",
        ("Projects", "city"): "Thane",
        ("Projects", "locality"): "Ghodbunder Road",
        ("Projects", "pincode"): "400607",
        ("Projects", "construction_status"): "UNDER_CONSTRUCTION",
        ("Towers", "project_id"): 1,
        ("Towers", "name"): "Tower A",
        ("Towers", "floors"): 20,
        ("Units", "project_id"): 1,
        ("Units", "tower_id"): 1,
        ("Units", "unit_number"): "A-101",
        ("Units", "floor"): 1,
        ("Units", "bhk"): "2 BHK",
        ("Units", "property_type"): "APARTMENT",
        ("Units", "carpet_area"): 850,
        ("Units", "price"): 9500000,
        ("Units", "availability"): "AVAILABLE",
        ("Units", "status"): "ACTIVE",
    }
    return examples.get((sheet, field), None)


def _template_workbook() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    instructions = [
        ["MetroCRM Import Template"],
        ["How to use"],
        ["1", "Keep the column names unchanged."],
        ["2", "Leave id blank to create a new record. Put an existing id to update that exact record."],
        ["3", "Import sheets in this workbook together; parent records must exist before child records."],
        ["4", "Do not change IDs when you want an update."],
        ["5", "JSON fields such as amenities_json must contain valid JSON (for example [\"Gym\",\"Pool\"])."],
        ["6", "Dates should use Excel date cells or ISO format YYYY-MM-DD / YYYY-MM-DD HH:MM:SS."],
        ["7", "The system validates all rows and rolls back the complete file if any row is invalid."],
        ["8", "Only Super Admin and Admin can import data."],
    ]
    for row in instructions:
        ws.append(row)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    ws["A1"].font = Font(bold=True, size=16)

    sheets = _models()
    for sheet, (model, fields) in sheets.items():
        sh = wb.create_sheet(sheet)
        sh.append(fields)
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
            cell.comment = Comment("Leave id blank for new record; use an existing id to update.", "MetroCRM") if cell.value == "id" else None
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for idx, field in enumerate(fields, start=1):
            sh.column_dimensions[get_column_letter(idx)].width = min(34, max(12, len(field) + 4))
        sh.sheet_view.showGridLines = False
    examples = wb.create_sheet("Examples")
    examples.append(["Sheet", "Field", "Example Value"])
    for cell in examples[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for sheet, (_, fields) in sheets.items():
        for field in fields:
            example = _sample_value(sheet, field)
            if example is not None:
                examples.append([sheet, field, example])
    examples.freeze_panes = "A2"
    examples.column_dimensions["A"].width = 26
    examples.column_dimensions["B"].width = 32
    examples.column_dimensions["C"].width = 55
    return _save_workbook(wb)


def _save_workbook(wb: Workbook) -> BytesIO:
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _excel_value(value: Any):
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@", "\t", "\r", "\n", "\x00", "＝", "＋", "－", "＠"}:
        return "'" + value
    return value


def _export_workbook(entity_names: list[str], db: Session, scope_ids: set[int] | None = None) -> BytesIO:
    from app.models import User
    wb = Workbook()
    wb.remove(wb.active)
    model_map = _models()
    scope_fields = {
        "Leads": "assigned_user_id",
        "Customers": None,
        "CustomerRequirements": None,
        "Developers": None,
        "Projects": None,
        "Towers": None,
        "Units": None,
    }
    for entity in entity_names:
        if entity not in model_map:
            continue
        model, fields = model_map[entity]
        sh = wb.create_sheet(entity)
        sh.append(fields)
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        query = select(model)
        owner_field = scope_fields.get(entity)
        if scope_ids and owner_field:
            query = query.where(getattr(model, owner_field).in_(scope_ids))
        rows = db.scalars(query.order_by(model.id.asc())).all()
        for row in rows:
            sh.append([_excel_value(getattr(row, field, None)) for field in fields])
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
    if not wb.sheetnames:
        wb.create_sheet("Export")
    return _save_workbook(wb)


def _export_report_workbook(actor, db: Session) -> BytesIO:
    data = _report_data(actor, db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Report Summary"
    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for key, value in data["metrics"].items():
        ws.append([key, _excel_value(value)])
    team = wb.create_sheet("Team Scope")
    team.append(["User ID", "Username", "Role", "Reporting Manager ID"])
    for user in data["users"]:
        team.append([user.id, user.username, user.role, user.reporting_manager_id])
    return _save_workbook(wb)


def _convert_excel_value(column, value):
    if value is None or value == "":
        return None
    from sqlalchemy import Boolean, DateTime, Integer, Numeric
    col_type = column.type
    if isinstance(col_type, Boolean):
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in {"1", "true", "yes", "y"}
    if isinstance(col_type, Integer):
        return int(value)
    if isinstance(col_type, Numeric):
        return Decimal(str(value))
    if isinstance(col_type, DateTime):
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    if isinstance(value, (list, dict)):
        return json.dumps(value)
    return str(value)


def _import_workbook(content: bytes, db: Session) -> tuple[int, int, list[str]]:
    model_map = _models()
    wb = load_workbook(BytesIO(content), data_only=True)
    updated = 0
    created = 0
    errors: list[str] = []
    try:
        for sheet_name in wb.sheetnames:
            if sheet_name not in model_map or sheet_name == "Instructions":
                continue
            model, allowed_fields = model_map[sheet_name]
            ws = wb[sheet_name]
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
            if not headers or headers[0] != "id":
                errors.append(f"{sheet_name}: first column must be id")
                continue
            field_set = set(allowed_fields)
            unknown = [h for h in headers if h and h not in field_set]
            if unknown:
                errors.append(f"{sheet_name}: unsupported columns: {', '.join(unknown)}")
                continue
            columns = {col.name: col for col in model.__table__.columns}
            for row_no, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(v not in (None, "") for v in row_values):
                    continue
                raw = {headers[i]: row_values[i] for i in range(min(len(headers), len(row_values))) if headers[i]}
                try:
                    row_id = int(raw["id"]) if raw.get("id") not in (None, "") else None
                    obj = db.get(model, row_id) if row_id else model()
                    if row_id and obj is None:
                        raise ValueError(f"id {row_id} does not exist")
                    for field, value in raw.items():
                        if field == "id" or field not in columns:
                            continue
                        converted = _convert_excel_value(columns[field], value)
                        if field.endswith("_json") and converted is not None:
                            json.loads(converted)
                        setattr(obj, field, converted)
                    db.add(obj)
                    if row_id:
                        updated += 1
                    else:
                        created += 1
                except Exception as exc:
                    errors.append(f"{sheet_name} row {row_no}: {exc}")
                    break
            if errors:
                break
        if errors:
            db.rollback()
        else:
            db.commit()
    except Exception as exc:
        db.rollback()
        errors.append(f"Workbook validation failed: {exc}")
    return created, updated, errors


def _descendant_ids(actor, db: Session) -> set[int]:
    if actor.role in {"SUPER_ADMIN", "ADMIN"}:
        return set(db.scalars(select(type(actor).id)).all())
    ids = {actor.id}
    changed = True
    while changed:
        changed = False
        children = db.scalars(select(type(actor)).where(type(actor).reporting_manager_id.in_(ids))).all()
        for child in children:
            if child.id not in ids:
                ids.add(child.id)
                changed = True
    return ids


def _report_scope(actor, db: Session) -> set[int]:
    return _descendant_ids(actor, db)


def _report_data(actor, db: Session) -> dict[str, Any]:
    from app.models import Lead, FollowUp, SiteVisit, Token, Booking, Commission, User
    ids = _report_scope(actor, db)
    leads = db.scalars(select(Lead).where(Lead.assigned_user_id.in_(ids))).all()
    followups = db.scalars(select(FollowUp).where(FollowUp.assigned_user_id.in_(ids))).all()
    visits = db.scalars(select(SiteVisit).where(or_(SiteVisit.sales_employee_id.in_(ids), SiteVisit.conducted_by_id.in_(ids)))).all()
    bookings = db.scalars(select(Booking).where(or_(Booking.employee_id.in_(ids), Booking.manager_id.in_(ids)))).all()
    tokens = db.scalars(select(Token).where(or_(Token.salesperson_id.in_(ids), Token.closing_manager_id.in_(ids)))).all()
    commissions = db.scalars(select(Commission).where(Commission.employee_id.in_(ids))).all()
    users = db.scalars(select(User).where(User.id.in_(ids)).order_by(User.id)).all()
    return {
        "ids": ids,
        "users": users,
        "leads": leads,
        "followups": followups,
        "visits": visits,
        "bookings": bookings,
        "tokens": tokens,
        "commissions": commissions,
        "metrics": {
            "leads": len(leads),
            "new_leads": sum(1 for x in leads if x.status == "NEW"),
            "closed_leads": sum(1 for x in leads if str(x.status).upper() in {"CLOSED", "WON", "CONVERTED"}),
            "open_followups": sum(1 for x in followups if not x.completed),
            "overdue_followups": sum(1 for x in followups if not x.completed and x.due_at < datetime.now(timezone.utc).replace(tzinfo=None)),
            "visits": len(visits),
            "bookings": len(bookings),
            "booking_value": sum((x.sale_value or 0) for x in bookings),
            "token_value": sum((x.token_amount or 0) for x in tokens),
            "commission_pending": sum((x.pending_amount or 0) for x in commissions),
        },
    }


def _html_page(title: str, body: str) -> HTMLResponse:
    html = f"""<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>MetroCRM | {title}</title><style>body{{font-family:Inter,Arial,sans-serif;background:#f5f7fb;color:#10213f;margin:0;padding:24px}}.wrap{{max-width:1180px;margin:auto}}.card{{background:#fff;border:1px solid #e5eaf2;border-radius:16px;padding:22px;margin-bottom:18px;box-shadow:0 5px 20px #0000000a}}h1{{margin:0 0 6px}}h2{{font-size:18px}}.muted{{color:#64748b;font-size:13px}}.grid{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}}.stat{{border:1px solid #e5eaf2;border-radius:12px;padding:15px}}.value{{font-size:24px;font-weight:800;margin-top:5px}}.btn{{display:inline-block;background:#2563eb;color:#fff;text-decoration:none;border:0;border-radius:9px;padding:10px 14px;font-weight:700;cursor:pointer}}.btn.secondary{{background:#0f766e}}.btn.warn{{background:#b45309}}.btn.danger{{background:#b91c1c}}.row{{display:flex;gap:10px;align-items:center;flex-wrap:wrap}}table{{width:100%;border-collapse:collapse;font-size:13px}}th,td{{padding:10px;border-bottom:1px solid #eef2f7;text-align:left}}th{{background:#f8fafc}}input,select,textarea{{padding:9px;border:1px solid #cbd5e1;border-radius:8px;width:100%;box-sizing:border-box}}label{{font-size:12px;font-weight:700;display:block;margin:10px 0 5px}}.two{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}.ok{{color:#166534;background:#dcfce7;padding:4px 8px;border-radius:999px;font-size:11px;font-weight:700}}.pending{{color:#92400e;background:#fef3c7;padding:4px 8px;border-radius:999px;font-size:11px;font-weight:700}}@media(max-width:800px){{.grid,.two{{grid-template-columns:1fr}}}}</style></head><body><div class='wrap'><div class='row' style='justify-content:space-between;margin-bottom:18px'><div><h1>{title}</h1><div class='muted'>MetroCRM controlled data operations</div></div><a class='btn secondary' href='/dashboard'>Dashboard</a></div>{body}</div></body></html>"""
    return HTMLResponse(html)


@router.get("/admin/data-tools", response_class=HTMLResponse)
def data_tools_page(request: Request, db: Session = Depends(get_db)):
    actor = _guard_admin(request, db)
    if not actor:
        return RedirectResponse("/dashboard", status_code=303)
    from app.models import DataRequest, User
    pending = db.scalars(select(DataRequest).where(DataRequest.action == "EXPORT", DataRequest.status.in_(["PENDING", "APPROVED"])).order_by(DataRequest.id.desc()).limit(50)).all()
    pending = [r for r in pending if r.requester_id == actor.id or actor.role == "SUPER_ADMIN"]
    csrf = csrf_token(request)
    body = "<div class='card'><h2>Data Import / Export</h2><p class='muted'>Only Super Admin and Admin have access. Admin exports require Super Admin approval; imports do not require approval.</p>"
    body += "<div class='grid'>"
    for entity in _models():
        body += f"<div class='stat'><strong>{entity}</strong><p class='muted'>Excel data operation</p><a class='btn' href='/admin/data-tools/export/{entity}'>Export</a></div>"
    body += "</div></div>"
    body += f"<div class='card'><h2>Import Excel</h2><p class='muted'>Download the marketplace-style template, fill it, then upload it. Existing IDs are updated; blank IDs create new records.</p><div class='row'><a class='btn secondary' href='/admin/data-tools/template.xlsx'>Download Sample Excel Template</a></div><form method='post' action='/admin/data-tools/import' enctype='multipart/form-data'><input type='hidden' name='csrf_token' value='{csrf}'><label>Excel file (.xlsx)</label><input type='file' name='file' accept='.xlsx' required><div style='margin-top:12px'><button class='btn' type='submit'>Validate & Import</button></div></form></div>"
    if actor.role == "ADMIN":
        body += f"<div class='card'><h2>Admin Export Approval</h2><form method='post' action='/admin/data-tools/export-request'><input type='hidden' name='csrf_token' value='{csrf}'><div class='two'><div><label>Data set</label><select name='entity_type'>" + "".join(f"<option value='{e}'>{e}</option>" for e in list(_models()) + ["REPORTS"]) + "</select></div><div><label>Reason</label><input name='reason' required minlength='5' placeholder='Why is this export required?'></div></div><div style='margin-top:12px'><button class='btn warn' type='submit'>Request Super Admin Approval</button></div></form></div>"
    if pending:
        body += "<div class='card'><h2>Export Requests</h2><table><tr><th>ID</th><th>Requester</th><th>Data</th><th>Status</th><th>Action</th></tr>"
        for r in pending:
            if actor.role == "SUPER_ADMIN":
                status = "<span class='ok'>APPROVED</span>" if r.status == "APPROVED" else "<span class='pending'>PENDING</span>"
                action = f"<a class='btn' href='/super-admin/data-export-requests'>Review</a>"
            else:
                css_class = "ok" if r.status == "APPROVED" else "pending"
                status = f"<span class='{css_class}'>{r.status}</span>"
                action = f"<a class='btn' href='/admin/data-tools/export/{r.entity_type}?request_id={r.id}'>Download</a>" if r.status == "APPROVED" else "Waiting for Super Admin"
            body += f"<tr><td>#{r.id}</td><td>{r.requester_id}</td><td>{r.entity_type}</td><td>{status}</td><td>{action}</td></tr>"
        body += "</table></div>"
    return _html_page("Data Operations", body)


@router.get("/admin/data-tools/template.xlsx")
def download_template(request: Request, db: Session = Depends(get_db)):
    if not _guard_admin(request, db):
        return RedirectResponse("/dashboard", status_code=303)
    return StreamingResponse(_template_workbook(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=MetroCRM_Data_Import_Template.xlsx"})


@router.post("/admin/data-tools/import", dependencies=[Depends(require_csrf)])
def import_data(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    actor = _guard_admin(request, db)
    if not actor:
        return RedirectResponse("/dashboard", status_code=303)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return _html_page("Import Error", "<div class='card'><h2>Only .xlsx files are supported.</h2><a class='btn' href='/admin/data-tools'>Back</a></div>")
    content = file.file.read()
    if len(content) > 10 * 1024 * 1024:
        return _html_page("Import Error", "<div class='card'><h2>File is too large. Maximum size is 10 MB.</h2></div>")
    if not is_zipfile(BytesIO(content)):
        return _html_page("Import Error", "<div class='card'><h2>The uploaded file is not a valid XLSX workbook.</h2></div>")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            uncompressed_size = sum(info.file_size for info in infos)
            if len(infos) > 5000 or uncompressed_size > 50 * 1024 * 1024:
                return _html_page("Import Error", "<div class='card'><h2>Workbook contents exceed the safe processing limits.</h2></div>")
    except BadZipFile:
        return _html_page("Import Error", "<div class='card'><h2>The uploaded file is not a valid XLSX workbook.</h2></div>")
    created, updated, errors = _import_workbook(content, db)
    if errors:
        return _html_page("Import Validation Failed", "<div class='card'><h2>No changes were saved.</h2><p class='muted'>The complete workbook was rolled back because validation failed.</p><ul>" + "".join(f"<li>{e}</li>" for e in errors[:30]) + "</ul><a class='btn' href='/admin/data-tools'>Back</a></div>")
    _audit(db, actor.id, "DATA_IMPORT_COMPLETED", "Import", file.filename, {"created": created, "updated": updated})
    db.commit()
    return _html_page("Import Complete", f"<div class='card'><h2>Import successful</h2><p>Created: <strong>{created}</strong> · Updated: <strong>{updated}</strong></p><a class='btn' href='/admin/data-tools'>Back to Data Operations</a></div>")


@router.post("/admin/data-tools/export-request", dependencies=[Depends(require_csrf)])
def request_export(request: Request, entity_type: str = Form(...), reason: str = Form(...), db: Session = Depends(get_db)):
    actor = _guard_admin(request, db)
    if not actor or actor.role != "ADMIN":
        return RedirectResponse("/admin/data-tools", status_code=303)
    allowed = set(_models()) | {"REPORTS"}
    if entity_type not in allowed or len(reason.strip()) < 5:
        return RedirectResponse("/admin/data-tools?error=validation", status_code=303)
    from app.models import DataRequest
    duplicate = db.scalar(select(DataRequest).where(DataRequest.requester_id == actor.id, DataRequest.action == "EXPORT", DataRequest.entity_type == entity_type, DataRequest.status.in_(["PENDING", "APPROVED"])).limit(1))
    if duplicate:
        return RedirectResponse("/admin/data-tools", status_code=303)
    row = DataRequest(requester_id=actor.id, action="EXPORT", entity_type=entity_type, filters_json="{}", reason=reason.strip(), status="PENDING")
    db.add(row)
    db.flush()
    _audit(db, actor.id, "DATA_EXPORT_REQUESTED", "DataRequest", str(row.id), {"entity_type": entity_type})
    db.commit()
    return RedirectResponse("/admin/data-tools", status_code=303)


@router.get("/super-admin/data-export-requests", response_class=HTMLResponse)
def export_requests(request: Request, db: Session = Depends(get_db)):
    actor = _guard_super_admin(request, db)
    if not actor:
        return RedirectResponse("/super-admin/login", status_code=303)
    from app.models import DataRequest, User
    rows = db.scalars(select(DataRequest).where(DataRequest.action == "EXPORT", DataRequest.status.in_(["PENDING", "APPROVED"])).order_by(DataRequest.id.desc()).limit(100)).all()
    csrf = csrf_token(request)
    body = "<div class='card'><h2>Admin Export Approval Requests</h2><table><tr><th>ID</th><th>Requester</th><th>Dataset</th><th>Reason</th><th>Status</th><th>Action</th></tr>"
    for row in rows:
        user = db.get(User, row.requester_id)
        status = "<span class='ok'>APPROVED</span>" if row.status == "APPROVED" else "<span class='pending'>PENDING</span>"
        action = ""
        if row.status == "PENDING":
            action = f"<form class='row' method='post' action='/super-admin/data-export-requests/{row.id}/approve'><input type='hidden' name='csrf_token' value='{csrf}'><button class='btn' type='submit'>Approve</button><button class='btn danger' name='decision' value='reject' type='submit'>Reject</button></form>"
        elif row.status == "APPROVED":
            action = f"<a class='btn secondary' href='/super-admin/data-export-requests/{row.id}/download'>Download</a>"
        body += f"<tr><td>#{row.id}</td><td>{escape(user.username) if user else row.requester_id}</td><td>{row.entity_type}</td><td>{escape(str(row.reason))}</td><td>{status}</td><td>{action}</td></tr>"
    body += "</table></div>"
    return _html_page("Export Approval", body)


@router.post("/super-admin/data-export-requests/{request_id}/approve", dependencies=[Depends(require_csrf)])
def decide_export_request(request_id: int, request: Request, decision: str = Form("approve"), db: Session = Depends(get_db)):
    actor = _guard_super_admin(request, db)
    if not actor:
        return RedirectResponse("/super-admin/login", status_code=303)
    from app.models import DataRequest
    row = db.get(DataRequest, request_id)
    if not row or row.action != "EXPORT" or row.status != "PENDING":
        return RedirectResponse("/super-admin/data-export-requests", status_code=303)
    row.status = "APPROVED" if decision != "reject" else "REJECTED"
    row.decided_by_id = actor.id
    row.decided_at = datetime.now(timezone.utc).replace(tzinfo=None)
    row.approval_reason = "Approved by Super Admin" if row.status == "APPROVED" else "Rejected by Super Admin"
    _audit(db, actor.id, "DATA_EXPORT_APPROVED" if row.status == "APPROVED" else "DATA_EXPORT_REJECTED", "DataRequest", str(row.id), {"requester_id": row.requester_id, "entity_type": row.entity_type})
    db.commit()
    return RedirectResponse("/super-admin/data-export-requests", status_code=303)


@router.get("/admin/data-tools/export/{entity_type}")
def export_data(request: Request, entity_type: str, request_id: int | None = None, db: Session = Depends(get_db)):
    actor = _guard_admin(request, db)
    if not actor:
        return RedirectResponse("/dashboard", status_code=303)
    if entity_type != "REPORTS" and entity_type not in _models():
        return _html_page("Export", "<div class='card'><h2>Unknown dataset.</h2></div>")
    from app.models import DataRequest
    if actor.role == "ADMIN":
        if not request_id:
            return _html_page("Export Approval Required", f"<div class='card'><h2>Super Admin approval required</h2><p>You cannot directly download <strong>{entity_type}</strong>. Submit an approval request from the Data Operations page.</p><a class='btn' href='/admin/data-tools'>Request Approval</a></div>")
        approval = db.get(DataRequest, request_id)
        if not approval or approval.requester_id != actor.id or approval.action != "EXPORT" or approval.entity_type != entity_type or approval.status != "APPROVED":
            return _html_page("Export Denied", "<div class='card'><h2>This export is not approved.</h2><a class='btn' href='/admin/data-tools'>Back</a></div>")
        workbook = _export_report_workbook(actor, db) if entity_type == "REPORTS" else _export_workbook([entity_type], db)
        approval.status = "COMPLETED"
        _audit(db, actor.id, "DATA_EXPORT_COMPLETED", "DataRequest", str(approval.id), {"entity_type": entity_type})
        db.commit()
    else:
        workbook = _export_report_workbook(actor, db) if entity_type == "REPORTS" else _export_workbook([entity_type], db)
        _audit(db, actor.id, "DATA_EXPORT_COMPLETED", entity_type, None, {"approval": "super_admin_direct"})
        db.commit()
    return StreamingResponse(workbook, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=MetroCRM_{entity_type}.xlsx"})


@router.get("/super-admin/data-export-requests/{request_id}/download")
def super_admin_download_export(request_id: int, request: Request, db: Session = Depends(get_db)):
    actor = _guard_super_admin(request, db)
    if not actor:
        return RedirectResponse("/super-admin/login", status_code=303)
    from app.models import DataRequest
    row = db.get(DataRequest, request_id)
    if not row or row.action != "EXPORT" or row.status != "APPROVED":
        return RedirectResponse("/super-admin/data-export-requests", status_code=303)
    workbook = _export_workbook([row.entity_type], db)
    row.status = "COMPLETED"
    _audit(db, actor.id, "DATA_EXPORT_COMPLETED", "DataRequest", str(row.id), {"entity_type": row.entity_type, "approved_requester_id": row.requester_id})
    db.commit()
    return StreamingResponse(workbook, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename=MetroCRM_{row.entity_type}_approved.xlsx"})


@router.get("/reports", response_class=HTMLResponse)
def reports_page(request: Request, db: Session = Depends(get_db)):
    actor = _current_user(request, db)
    if not actor:
        return RedirectResponse("/", status_code=303)
    data = _report_data(actor, db)
    m = data["metrics"]
    cards = [("Leads", m["leads"]), ("New Leads", m["new_leads"]), ("Closed / Won", m["closed_leads"]), ("Open Follow-ups", m["open_followups"]), ("Overdue Follow-ups", m["overdue_followups"]), ("Site Visits", m["visits"]), ("Bookings", m["bookings"]), ("Booking Value", f"₹ {float(m['booking_value']):,.2f}"), ("Token Value", f"₹ {float(m['token_value']):,.2f}"), ("Pending Commission", f"₹ {float(m['commission_pending']):,.2f}")]
    body = "<div class='card'><h2>Role-scoped Reports</h2><p class='muted'>Your role controls the data scope: Super Admin/Admin = all data; Manager/TL = self + reporting tree; other roles = self only.</p><div class='grid'>" + "".join(f"<div class='stat'><div class='muted'>{label}</div><div class='value'>{value}</div></div>" for label, value in cards) + "</div></div>"
    body += "<div class='card'><h2>Team Scope</h2><table><tr><th>User</th><th>Role</th><th>Reporting Manager</th></tr>"
    for user in data["users"]:
        manager = db.get(type(actor), user.reporting_manager_id) if user.reporting_manager_id else None
        body += f"<tr><td>{escape(user.username)}</td><td>{escape(str(user.role))}</td><td>{escape(manager.username) if manager else '—'}</td></tr>"
    body += "</table></div>"
    if actor.role == "SUPER_ADMIN":
        body += "<div class='card'><h2>Business Health</h2><p class='muted'>Super Admin only.</p><a class='btn secondary' href='/super-admin/business-health'>Open Business Health Report</a></div>"
    return _html_page("Reports", body)


@router.get("/super-admin/business-health", response_class=HTMLResponse)
def business_health(request: Request, db: Session = Depends(get_db)):
    actor = _guard_super_admin(request, db)
    if not actor:
        return RedirectResponse("/super-admin/login", status_code=303)
    from app.models import Lead, FollowUp, SiteVisit, Booking, Unit, User, Commission
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    total_leads = db.scalar(select(func.count(Lead.id))) or 0
    open_leads = db.scalar(select(func.count(Lead.id)).where(Lead.status.not_in(["CLOSED", "WON", "CONVERTED"]))) or 0
    won_leads = db.scalar(select(func.count(Lead.id)).where(Lead.status.in_(["CLOSED", "WON", "CONVERTED"]))) or 0
    total_users = db.scalar(select(func.count(User.id)).where(User.is_active.is_(True))) or 0
    overdue = db.scalar(select(func.count(FollowUp.id)).where(FollowUp.completed.is_(False), FollowUp.due_at < now)) or 0
    visits = db.scalar(select(func.count(SiteVisit.id))) or 0
    bookings = db.scalar(select(func.count(Booking.id))) or 0
    sales = db.scalar(select(func.coalesce(func.sum(Booking.sale_value), 0))) or 0
    inventory = db.scalar(select(func.count(Unit.id)).where(Unit.status.not_in(["SOLD", "CANCELLED"]))) or 0
    available = db.scalar(select(func.count(Unit.id)).where(or_(Unit.availability == "AVAILABLE", Unit.status == "AVAILABLE"))) or 0
    commission_pending = db.scalar(select(func.coalesce(func.sum(Commission.pending_amount), 0))) or 0
    conversion = (won_leads / total_leads * 100) if total_leads else 0
    body = "<div class='card'><h2>Business Health Dashboard</h2><p class='muted'>Restricted to Super Admin. Use this as the executive operating view of the CRM.</p><div class='grid'>"
    metrics = [("Active Users", total_users), ("Total Leads", total_leads), ("Open Leads", open_leads), ("Lead Conversion", f"{conversion:.1f}%"), ("Overdue Follow-ups", overdue), ("Site Visits", visits), ("Bookings", bookings), ("Sales Value", f"₹ {float(sales):,.2f}"), ("Property Inventory", inventory), ("Available Units", available), ("Pending Commission", f"₹ {float(commission_pending):,.2f}")]
    body += "".join(f"<div class='stat'><div class='muted'>{label}</div><div class='value'>{value}</div></div>" for label, value in metrics) + "</div></div>"
    body += "<div class='card'><h2>Recommended management checks</h2><ul><li>Follow up overdue leads first.</li><li>Compare new-lead inflow against conversion and booking value.</li><li>Track available inventory against current demand.</li><li>Review pending commission and booking collections.</li><li>Review inactive users and unusual export/import activity in the audit log.</li></ul></div>"
    return _html_page("Business Health", body)
